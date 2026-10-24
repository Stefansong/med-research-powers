"""Tests for skills/manuscript-export/scripts/export_docx.py.

Run from the repo root:  python3 -m pytest tests/test_export_docx.py -q
Requires: python-docx, pyyaml (same as the script itself).
"""
import importlib.util
import os
import subprocess
import sys
from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("yaml")
from docx import Document  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills" / "manuscript-export" / "scripts" / "export_docx.py"
YAML = ROOT / "skills" / "manuscript-writing" / "references" / "journal-templates.yaml"

MINIMAL = {
    "title-page.md": "# Deep learning for bladder cancer on CT\n\n**Authors:** A. Author^1^\n",
    "abstract.md": ("# Abstract\n\n**Background:** Detection is hard.\n\n**Methods:** We trained a CNN.\n\n"
                    "**Results:** Sensitivity ~90%.\n\n**Conclusions:** Promising.\n"),
    "key-points.md": "# Key Points\n\n**Question:** Does it work?\n\n**Findings:** Yes.\n\n**Meaning:** Useful.\n",
    "research-in-context.md": ("# Research in Context\n\n**Evidence before this study:** Little.\n\n"
                               "**Added value of this study:** Some.\n\n**Implications:** More trials.\n"),
    "introduction.md": ("# Introduction\n\nPrior work reported ~90% sensitivity using 2^10 images.[1]\n\n"
                        "<!-- PLACEHOLDER: add funding sentence -->\n"),
    "methods.md": "# Methods\n\nWe included adults with *suspected* cancer; H~2~O was **not** used.\n",
    "results.md": ("# Results\n\n**Table 1.** Baseline characteristics\n\n"
                   "| Variable | Cases | Controls |\n|---|---|---|\n| Age | 65 | 63 |\n\n"
                   "**Figure 1.** ROC curve.\n"),
    "discussion.md": "# Discussion\n\nOur model achieved high sensitivity.\n",
    "references.md": ("# References\n\n1. Smith J. Bladder cancer on CT. *Radiology*. 2020;295:1-9.\n"
                      "2. Doe A. Deep learning in urology. Eur Urol. 2021;80:10-20.\n"),
    "supplementary.md": "# Supplementary\n\n**eTable 1.** Extra\n\n| A | B |\n|---|---|\n| 1 | 2 |\n",
}


@pytest.fixture
def manuscript(tmp_path):
    d = tmp_path / "manuscript"
    d.mkdir()
    for name, content in MINIMAL.items():
        (d / name).write_text(content, encoding="utf-8")
    return d


def run_export(manuscript_dir, journal, out, *extra, env=None, cwd=None):
    cmd = [sys.executable, str(SCRIPT), "--manuscript-dir", str(manuscript_dir), "--journal", journal,
           "--output", str(out), "--yaml", str(YAML), *extra]
    full_env = {**os.environ, **(env or {})}
    return subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace",
                          cwd=str(cwd or manuscript_dir.parent), env=full_env)


def headings(docx_path):
    return [p.text for p in Document(str(docx_path)).paragraphs if p.style.name.startswith("Heading")]


def full_text(docx_path):
    return "\n".join(p.text for p in Document(str(docx_path)).paragraphs)


@pytest.mark.parametrize("journal", ["european-urology", "nature", "lancet", "jama-network-open"])
def test_export_exits_zero_and_writes_docx(manuscript, tmp_path, journal):
    out = tmp_path / f"{journal}.docx"
    r = run_export(manuscript, journal, out, "--report", str(tmp_path / f"{journal}-report.md"))
    assert r.returncode == 0, r.stdout + r.stderr
    assert out.exists() and out.stat().st_size > 0
    assert (tmp_path / f"{journal}-report.md").exists()


def test_nature_order_is_i_r_d_m(manuscript, tmp_path):
    out = tmp_path / "nature.docx"
    assert run_export(manuscript, "nature", out).returncode == 0
    h = headings(out)
    idx = [h.index(x) for x in ("Introduction", "Results", "Discussion", "Methods")]
    assert idx == sorted(idx), h


def test_standard_order_is_i_m_r_d(manuscript, tmp_path):
    out = tmp_path / "eu.docx"
    assert run_export(manuscript, "european-urology", out).returncode == 0
    h = headings(out)
    idx = [h.index(x) for x in ("Introduction", "Methods", "Results", "Discussion")]
    assert idx == sorted(idx), h


def test_lancet_includes_research_in_context(manuscript, tmp_path):
    out = tmp_path / "lancet.docx"
    r = run_export(manuscript, "lancet", out)
    assert r.returncode == 0
    h = headings(out)
    assert "Research in Context" in h
    assert h.index("Research in Context") < h.index("Introduction")
    assert "Key Points" not in h  # not part of the Lancet family order
    assert "research-in-context" in r.stdout


def test_jama_includes_key_points(manuscript, tmp_path):
    out = tmp_path / "jama.docx"
    r = run_export(manuscript, "jama-network-open", out)
    assert r.returncode == 0
    h = headings(out)
    assert "Key Points" in h
    assert h.index("Key Points") < h.index("Abstract")
    assert "Research in Context" not in h


def test_unpaired_markers_survive_in_body(manuscript, tmp_path):
    out = tmp_path / "eu.docx"
    assert run_export(manuscript, "european-urology", out).returncode == 0
    text = full_text(out)
    assert "~90%" in text
    assert "2^10" in text
    assert "**" not in text and "~2~" not in text  # closed pairs were converted


def test_placeholder_is_reported(manuscript, tmp_path):
    report = tmp_path / "report.md"
    r = run_export(manuscript, "european-urology", tmp_path / "eu.docx", "--report", str(report))
    assert r.returncode == 0
    txt = report.read_text(encoding="utf-8")
    assert "PLACEHOLDER" in txt and "⚠️" in txt
    assert "introduction.md:" in txt


def test_missing_family_section_is_flagged(manuscript, tmp_path):
    (manuscript / "key-points.md").unlink()
    r = run_export(manuscript, "jama-network-open", tmp_path / "jama.docx")
    assert r.returncode == 0
    assert "⚠️" in r.stdout and "key-points.md" in r.stdout


def test_word_buckets_are_separate(manuscript, tmp_path):
    r = run_export(manuscript, "european-urology", tmp_path / "eu.docx")
    assert r.returncode == 0
    assert "| Abstract |" in r.stdout
    assert "| References (entries) | 2 |" in r.stdout
    body_line = next(l for l in r.stdout.splitlines() if l.startswith("| Body ("))
    body_words = int(body_line.split("|")[2].strip())
    assert 0 < body_words < 40  # abstract / title page / references are not in the body bucket


def test_unknown_journal_exits_2(manuscript, tmp_path):
    r = run_export(manuscript, "no-such-journal-xyz", tmp_path / "x.docx")
    assert r.returncode == 2
    assert "--search" in r.stderr
    assert not (tmp_path / "x.docx").exists()


def test_supplementary_flag_writes_second_docx(manuscript, tmp_path):
    out = tmp_path / "manuscript" / "manuscript.docx"
    r = run_export(manuscript, "european-urology", out, "--supplementary")
    assert r.returncode == 0
    assert (tmp_path / "manuscript" / "supplementary.docx").exists()
    assert "supplementary.docx" in r.stdout


def test_runs_under_c_locale(manuscript, tmp_path):
    env = {"LC_ALL": "C", "LANG": "C", "PYTHONUTF8": "0", "PYTHONCOERCECLOCALE": "0", "PYTHONIOENCODING": ""}
    r = run_export(manuscript, "lancet", tmp_path / "lancet.docx", "--report", str(tmp_path / "r.md"), env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "Research in Context" in headings(tmp_path / "lancet.docx")


# ─── unit tests of the inline formatter ──────────────────────────────────────

def _load_module():
    spec = importlib.util.spec_from_file_location("export_docx", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _runs(mod, text):
    doc = Document()
    p = doc.add_paragraph()
    mod.add_formatted_text(p, text, {"font": "Times New Roman"})
    return [(r.text, bool(r.bold), bool(r.italic), bool(r.font.superscript), bool(r.font.subscript))
            for r in p.runs]


def test_inline_formatting_keeps_unpaired_markers():
    mod = _load_module()
    assert "".join(t for t, *_ in _runs(mod, "~90% sensitivity")) == "~90% sensitivity"
    assert "".join(t for t, *_ in _runs(mod, "2^10")) == "2^10"
    assert "".join(t for t, *_ in _runs(mod, "marked with * were")) == "marked with * were"
    assert mod.plain_text("~90% and 2^10") == "~90% and 2^10"


def test_inline_formatting_converts_closed_pairs():
    mod = _load_module()
    assert ("bold", True, False, False, False) in _runs(mod, "**bold**")
    assert ("it", False, True, False, False) in _runs(mod, "*it*")
    assert ("2", False, False, True, False) in _runs(mod, "10^2^")
    assert ("2", False, False, False, True) in _runs(mod, "H~2~O")
    assert "".join(t for t, *_ in _runs(mod, "H~2~O")) == "H2O"


@pytest.mark.parametrize("value,limit", [
    ("≤40, Vancouver", 40),
    ("150-250 words", 250),          # a range: the upper bound, not the first number
    ("200–300 words", 300),
    ("2,500 words", 2500),
    ("8 pages", None),               # a page limit is not a word limit (ieee-jbhi)
    ("≤10 pages", None),             # ieee-tmi
    ("unlimited", None),
])
def test_parse_limit(value, limit):
    spec = importlib.util.spec_from_file_location("export_docx", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    assert module._parse_limit(value) == limit


@pytest.mark.parametrize("value", [
    "no hard limit (~25 per 1,000 words; up to ~150 for Reviews), Nature",
    "no fixed limit (figure legends ≤350 words each)",
    "no formal limit",
    "no strict limit",
    "8–10 printed pages including figures, tables and references (~650 words per page)",
    "≤10 double-column pages",
])
def test_parse_limit_no_limit_and_page_phrases(value):
    assert _load_module()._parse_limit(value) is None


@pytest.mark.parametrize("value,limit", [
    ("≤100 (recommended), AMA", 100),     # soft limits still give a number (a ⚠️ at most)
    ("flexible (usually <100)", 100),
    ("3000 words, 6 pages", 3000),        # the page count belongs to another number
])
def test_parse_limit_soft_limits_keep_their_number(value, limit):
    assert _load_module()._parse_limit(value) == limit


# ─── malformed journal-overrides.yaml ────────────────────────────────────────

@pytest.mark.parametrize("content", [
    "- just a string\n- another\n",                  # a list of strings
    "journals:\n  - id: x\n",                        # a mapping without `templates`
    "just some text\n",                              # a scalar
    "templates:\n  id: x\n",                         # `templates` is not a list
    "templates:\n  - journal: No Id Journal\n",      # an entry without `id`
    "templates: [\n",                                # invalid YAML
])
def test_malformed_overrides_exit_2_with_one_error_line(manuscript, tmp_path, content):
    ov = tmp_path / "bad-overrides.yaml"
    ov.write_text(content, encoding="utf-8")
    out = tmp_path / "x.docx"
    r = run_export(manuscript, "european-urology", out, "--overrides", str(ov))
    assert r.returncode == 2, r.stdout + r.stderr
    assert "Traceback" not in r.stderr
    errors = [l for l in r.stderr.splitlines() if l.startswith("Error:")]
    assert len(errors) == 1 and errors[0].startswith("Error: overrides file:"), r.stderr
    assert not out.exists()


def test_malformed_default_overrides_in_cwd_is_an_error(manuscript, tmp_path):
    (tmp_path / "journal-overrides.yaml").write_text("templates: 42\n", encoding="utf-8")  # cwd = tmp_path
    r = run_export(manuscript, "european-urology", tmp_path / "x.docx")
    assert r.returncode == 2 and "Error: overrides file:" in r.stderr and "Traceback" not in r.stderr


def test_empty_overrides_file_is_accepted(manuscript, tmp_path):
    ov = tmp_path / "empty.yaml"
    ov.write_text("", encoding="utf-8")
    r = run_export(manuscript, "european-urology", tmp_path / "x.docx", "--overrides", str(ov))
    assert r.returncode == 0, r.stderr


def test_missing_named_overrides_warns_and_uses_library(manuscript, tmp_path):
    r = run_export(manuscript, "european-urology", tmp_path / "x.docx", "--overrides", str(tmp_path / "nope.yaml"))
    assert r.returncode == 0, r.stderr
    assert "Warning: overrides file not found" in r.stderr


def test_valid_overrides_entry_is_used(manuscript, tmp_path):
    ov = tmp_path / "ov.yaml"
    ov.write_text("templates:\n  - id: my-journal\n    journal: My Journal\n    family: jama\n"
                  "    word_limit: 3000 words\n", encoding="utf-8")
    out = tmp_path / "mine.docx"
    r = run_export(manuscript, "my-journal", out, "--overrides", str(ov))
    assert r.returncode == 0, r.stderr
    assert "Key Points" in headings(out)


# ─── Markdown block parsing (the .docx is opened and inspected) ──────────────

@pytest.fixture(scope="module")
def mod():
    return _load_module()


def export_md(mod, tmp_path, md, section="results.md", journal="european-urology"):
    """Export one section file; returns (Document, stats, report text)."""
    d = tmp_path / "ms"
    d.mkdir(exist_ok=True)
    (d / section).write_text(md, encoding="utf-8")
    config = mod.get_journal_config(journal, str(YAML))
    out = tmp_path / "out.docx"
    stats = mod.build_docx(str(d), config, str(out))
    return Document(str(out)), stats, mod.generate_report(stats, config, str(out))


def paras(doc):
    return [(p.style.name, p.text) for p in doc.paragraphs]


def texts(doc):
    return [p.text for p in doc.paragraphs if p.text]


def test_soft_wrapped_lines_join_into_one_paragraph(mod, tmp_path):
    md = ("# Results\n\nThe first line of a paragraph\nthat was soft-wrapped\n   over three lines.\n\n"
          "Second paragraph.\n")
    doc, stats, _ = export_md(mod, tmp_path, md)
    assert texts(doc) == ["Results", "The first line of a paragraph that was soft-wrapped over three lines.",
                          "Second paragraph."]
    assert stats["words"]["results"] == 14


def test_cjk_lines_join_without_space(mod, tmp_path):
    md = "中文第一行\n第二行。\n\n结果显示，\nAUC 为 0.9\n\nEnglish words\n中文\n\n**结果**\n显示良好\n"
    doc, _, _ = export_md(mod, tmp_path, md)
    assert texts(doc) == ["中文第一行第二行。", "结果显示，AUC 为 0.9", "English words 中文", "结果显示良好"]


def test_hard_line_break_is_kept(mod, tmp_path):
    doc, _, _ = export_md(mod, tmp_path, "Line one  \nLine two\\\nLine three\n")
    assert texts(doc) == ["Line one\nLine two\nLine three"]


def test_bold_label_lines_joined_are_reported(mod, tmp_path):
    doc, stats, report = export_md(mod, tmp_path, "**Background:** Hard.\n**Methods:** A CNN.\n", "abstract.md")
    assert texts(doc) == ["Background: Hard. Methods: A CNN."]
    assert any(w.startswith("abstract.md:1 — lines 1–2 were joined") for w in stats["warnings"])
    assert "put a blank line between them" in report


def test_fenced_code_block_is_verbatim(mod, tmp_path):
    md = ("Intro text.\n\n```python\n# not a heading\nx = **not bold** * 2\n    indented\n"
          "<!-- kept inside code -->\n```\n\n~~~\n## also code\n- not a bullet\n~~~\n")
    doc, stats, _ = export_md(mod, tmp_path, md)
    assert not any(s.startswith("Heading") or s.startswith("List") for s, _ in paras(doc))
    code = [p for p in doc.paragraphs if p.text.startswith("# not a heading")]
    assert len(code) == 1
    assert code[0].text == "# not a heading\nx = **not bold** * 2\n    indented\n<!-- kept inside code -->"
    assert all(r.font.name == "Courier New" and not r.bold for r in code[0].runs)
    assert "## also code\n- not a bullet" in texts(doc)
    assert stats["words"]["results"] == 2          # code is not counted as words


def test_unclosed_fence_is_reported(mod, tmp_path):
    doc, stats, report = export_md(mod, tmp_path, "Text.\n\n```\ncode to the end\n")
    assert "code to the end" in texts(doc)
    assert any("never closed" in w for w in stats["warnings"]) and "never closed" in report


def test_strip_html_comments_skips_code_fences(mod):
    raw = "a <!-- x -->b\n```\n<!-- keep -->\n```\n<!--\nmulti\n-->c\n"
    out = mod.strip_html_comments(raw)
    assert out.split("\n") == ["a b", "```", "<!-- keep -->", "```", "", "", "c", ""]


def test_heading_levels_4_to_6(mod, tmp_path):
    doc, _, _ = export_md(mod, tmp_path, "#### Level four\n\n##### Level five ##\n\n###### Level six\n")
    assert paras(doc) == [("Heading 4", "Level four"), ("Heading 5", "Level five"), ("Heading 6", "Level six")]


def test_horizontal_rules_are_not_text(mod, tmp_path):
    doc, _, _ = export_md(mod, tmp_path, "Above.\n\n---\n\n***\n\n_ _ _\n\nBelow.\n")
    assert texts(doc) == ["Above.", "Below."]
    rules = [p for p in doc.paragraphs if not p.text]
    assert len(rules) == 3 and all("w:pBdr" in p._p.xml for p in rules)


def test_setext_headings(mod, tmp_path):
    doc, _, _ = export_md(mod, tmp_path, "Main title\n===\n\nSub title\n---\n\nText.\n")
    assert paras(doc) == [("Heading 1", "Main title"), ("Heading 2", "Sub title"), ("Normal", "Text.")]


def test_nested_lists_keep_level(mod, tmp_path):
    md = ("- top item that\n  wraps onto a second line\n  - second level\n    - third level\n"
          "      - fourth level\n- back to top\n\n"
          "1. one\n   1. sub-one\n   - sub-bullet\n2. two\n\n   Continuation paragraph of two.\n")
    doc, stats, _ = export_md(mod, tmp_path, md)
    assert paras(doc) == [
        ("List Bullet", "top item that wraps onto a second line"),
        ("List Bullet 2", "second level"),
        ("List Bullet 3", "third level"),
        ("List Bullet 3", "fourth level"),       # deepest style the template has
        ("List Bullet", "back to top"),
        ("List Number", "1. one"),
        ("List Number 2", "1. sub-one"),
        ("List Bullet 2", "sub-bullet"),
        ("List Number", "2. two"),
        ("List Continue", "Continuation paragraph of two."),
    ]
    numbered = [p for p in doc.paragraphs if p.style.name.startswith("List Number")]
    assert all('w:numId w:val="0"' in p._p.xml for p in numbered)     # literal numbers, no auto-numbering
    assert numbered[1].paragraph_format.left_indent > numbered[0].paragraph_format.left_indent


def test_nested_reference_items_are_not_counted(mod, tmp_path):
    md = "# References\n\n1. Smith J. A. 2020.\n   1. erratum note\n2. Doe A. B. 2021.\n"
    _, stats, _ = export_md(mod, tmp_path, md, "references.md")
    assert stats["ref_count"] == 2


def test_table_alignment_row_and_escaped_pipe(mod, tmp_path):
    md = "| Centre | Left | Right |\n|:-:|:--|--:|\n| 1 | a \\| b | 3 |\n"
    doc, stats, _ = export_md(mod, tmp_path, md)
    (t,) = doc.tables
    assert [[c.text for c in r.cells] for r in t.rows] == [["Centre", "Left", "Right"], ["1", "a | b", "3"]]
    assert [c.paragraphs[0].alignment for c in t.rows[1].cells] == [
        mod.WD_ALIGN_PARAGRAPH.CENTER, mod.WD_ALIGN_PARAGRAPH.LEFT, mod.WD_ALIGN_PARAGRAPH.RIGHT]
    assert stats["table_count"] == 1 and not stats["warnings"]


def test_table_without_outer_pipes(mod, tmp_path):
    md = "Caption line\nGroup | *n* | Mean\n--- | --- | ---\nA | 10 | 1.5\nB | 12 | 2.5\n\nAfter.\n"
    doc, stats, _ = export_md(mod, tmp_path, md)
    (t,) = doc.tables
    assert [[c.text for c in r.cells] for r in t.rows] == [
        ["Group", "n", "Mean"], ["A", "10", "1.5"], ["B", "12", "2.5"]]
    assert texts(doc) == ["Caption line", "After."]


def test_table_extra_cells_are_kept_and_reported(mod, tmp_path):
    md = "| A | B |\n|---|---|\n| 1 | 2 | 3 |\n| 4 | 5 | |\n"
    doc, stats, report = export_md(mod, tmp_path, md)
    (t,) = doc.tables
    assert [[c.text for c in r.cells] for r in t.rows] == [["A", "B", ""], ["1", "2", "3"], ["4", "5", ""]]
    assert len(stats["warnings"]) == 1 and stats["warnings"][0].startswith("results.md:3 — table row has 3 cells")
    assert "results.md:3 — table row has 3 cells but the header has 2" in report


def test_inline_nested_emphasis(mod):
    assert _runs(mod, "***x***") == [("x", True, True, False, False)]
    assert _runs(mod, "**a *b* c**") == [("a ", True, False, False, False), ("b", True, True, False, False),
                                         (" c", True, False, False, False)]
    assert _runs(mod, "*a **b** c*") == [("a ", False, True, False, False), ("b", True, True, False, False),
                                         (" c", False, True, False, False)]
    assert mod.plain_text("***x*** **a *b* c**") == "x a b c"
    assert _runs(mod, "**方法：**本研究")[0] == ("方法：", True, False, False, False)
    assert mod.plain_text("Age*, BMI* and 0.82***") == "Age*, BMI* and 0.82***"


def test_underscore_emphasis_respects_word_boundaries(mod):
    runs = _runs(mod, "_italic_ and __bold__ but snake_case_names, file_1_v2.csv and __init__.py")
    assert ("italic", False, True, False, False) in runs
    assert ("bold", True, False, False, False) in runs
    text = "".join(t for t, *_ in runs)
    assert "snake_case_names" in text and "file_1_v2.csv" in text


def test_inline_code_escapes_and_html_tags(mod):
    doc = Document()
    p = doc.add_paragraph()
    mod.add_formatted_text(p, r"Use `a*b*c` with \*literal\* stars, x<sup>2</sup>, ~~old~~", {"font": "Arial"})
    assert p.text == "Use a*b*c with *literal* stars, x2, old"
    code = [r for r in p.runs if r.text == "a*b*c"]
    assert code and code[0].font.name == "Courier New"
    assert any(r.text == "2" and r.font.superscript for r in p.runs)
    assert any(r.text == "old" and r.font.strike for r in p.runs)


def test_links_become_hyperlinks(mod):
    doc = Document()
    p = doc.add_paragraph()
    mod.add_formatted_text(p, "See [the **protocol**](https://example.org/p?a=1) and <https://doi.org/10.1/x>; "
                              "![Figure 1](fig1.png) stays.", {"font": "Arial"})
    assert p.text == "See the protocol and https://doi.org/10.1/x; ![Figure 1](fig1.png) stays."
    assert [(h.text, h.url) for h in p.hyperlinks] == [("the protocol", "https://example.org/p?a=1"),
                                                         ("https://doi.org/10.1/x", "https://doi.org/10.1/x")]
    assert any(r.bold for r in p.hyperlinks[0].runs)
    assert mod.plain_text("[text](http://x.org)") == "text"


def test_pandoc_citations_are_kept_and_reported(mod, tmp_path):
    md = ("Prior work [@smith2020; @doe2021, p. 3] and @lee2019 disagree.\n"
          "Contact a.b@hospital.org or see [the site](https://x.org/@user).\n\n"
          "| Study | Ref |\n|---|---|\n| X | [@smith2020] |\n")
    doc, stats, report = export_md(mod, tmp_path, md)
    assert "[@smith2020; @doe2021, p. 3] and @lee2019" in texts(doc)[0]
    assert sorted({k for k, _ in stats["citations"]}) == ["doe2021", "lee2019", "smith2020"]
    assert "Citation keys not converted (3)" in report
    assert "`@smith2020` — results.md:1, results.md:6" in report
    assert "@hospital" not in report and "@user" not in report


def test_no_conversion_warnings_when_clean(manuscript, tmp_path):
    report = tmp_path / "r.md"
    assert run_export(manuscript, "european-urology", tmp_path / "eu.docx", "--report", str(report)).returncode == 0
    notes = report.read_text(encoding="utf-8").split("## Conversion notes\n")[1].split("\n## ")[0].strip()
    # the only note: H~2~O in methods.md was read as a Pandoc subscript (and that is said)
    assert len(notes.splitlines()) == 1 and "`~2~`" in notes
    assert notes.startswith("- ℹ️ methods.md:3 — `~x~` set as subscript")


def test_nothing_to_note_when_nothing_was_interpreted(mod, tmp_path):
    _, _, report = export_md(mod, tmp_path, "# Results\n\nPlain text, ~90% and 2^10.\n")
    assert "## Conversion notes\n- ✅ Nothing to note" in report


@pytest.mark.parametrize("snippet", [
    "Median follow-up was [待补：来源] months.",
    "AUC ［待补］ in the test set.",
    "AUC 【待补】 in the test set.",
    "Sites: [待填].",
    "| Metric | Value |\n|---|---|\n| AUC | 【待补】 |",
    "<!-- 待补：基金号 -->",
])
def test_chinese_placeholders_are_reported(mod, tmp_path, snippet):
    _, stats, report = export_md(mod, tmp_path, f"# Results\n\n{snippet}\n")
    assert stats["placeholders"] and stats["placeholders"][0].startswith("results.md:")
    assert "待" in stats["placeholders"][0]
    assert "No placeholders detected" not in report


# ─── journal limits: structured `limits` block first, conservative free text otherwise (D1) ──

import yaml  # noqa: E402

LIBRARY = {e["id"]: e for e in yaml.safe_load(YAML.read_text(encoding="utf-8"))["templates"]}
FREE_TEXT_FIELDS = ("journal", "family", "word_limit", "abstract", "references", "figures", "tables")


def free_text_entry(journal_id, **extra):
    """The library entry with only its free-text fields (as in a user's journal-overrides.yaml)."""
    e = {"id": f"{journal_id}-ft"}
    e.update({k: LIBRARY[journal_id][k] for k in FREE_TEXT_FIELDS if k in LIBRARY[journal_id]})
    e.update(extra)
    return e


def write_overrides(tmp_path, *entries):
    ov = tmp_path / "ov.yaml"
    ov.write_text(yaml.safe_dump({"templates": list(entries)}, allow_unicode=True), encoding="utf-8")
    return ov


def sane(config, journal_id):
    """The checks that must hold for these real entries whether the limits come from the
    structured block or from the free text."""
    lim = config["limits"]
    if journal_id in ("ieee-tmi", "ieee-jbhi"):          # page limits: never a 10-word body limit
        assert lim["body_words"] is None and lim["pages"] and 8 <= lim["pages"] <= 20, lim
        return
    assert lim["body_words"] is None or lim["body_words"] >= 2000, lim   # nature: never 6
    assert "methods" in lim["body_excludes"], lim      # all three limits exclude (STAR) Methods
    if journal_id in ("nature-medicine", "cancer-cell"):
        assert lim["body_words"] == {"nature-medicine": 4000, "cancer-cell": 7000}[journal_id]
        assert "references" in lim["body_excludes"]


@pytest.mark.parametrize("journal_id", ["nature", "nature-medicine", "cancer-cell", "ieee-tmi", "ieee-jbhi"])
def test_real_library_entries_give_sane_limits(mod, journal_id):
    sane(mod.get_journal_config(journal_id, str(YAML)), journal_id)


@pytest.mark.parametrize("journal_id", ["nature", "nature-medicine", "cancer-cell", "ieee-tmi", "ieee-jbhi"])
def test_free_text_of_real_entries_gives_sane_limits(mod, tmp_path, journal_id):
    ov = write_overrides(tmp_path, free_text_entry(journal_id))
    config = mod.get_journal_config(f"{journal_id}-ft", str(YAML), str(ov))
    assert config["limits_source"].startswith("free-text")
    sane(config, journal_id)
    if journal_id == "nature":                         # "typically ~2500 … to ~4300": not one number
        assert config["limits"]["body_words"] is None
        assert any("typical values" in n for n in config["limit_notes"])
    if journal_id == "cancer-cell":                    # "including main figure legends"
        assert "figure_legends" not in config["limits"]["body_excludes"]


SYNTH = {"introduction.md": "# Introduction\n\n" + " ".join(["intro"] * 500) + "\n",
         "methods.md": "# Methods\n\n" + " ".join(["method"] * 3000) + "\n",
         "results.md": "# Results\n\n" + " ".join(["result"] * 1500) + "\n",
         "discussion.md": "# Discussion\n\n" + " ".join(["discuss"] * 1000) + "\n"}


def synth_manuscript(tmp_path, files=SYNTH):
    d = tmp_path / "manuscript"
    d.mkdir(parents=True, exist_ok=True)
    for name, content in files.items():
        (d / name).write_text(content, encoding="utf-8")
    return d


def report_rows(stdout):
    return {l.split("|")[1].strip().split(" (")[0]: [c.strip() for c in l.split("|")[2:5]]
            for l in stdout.splitlines() if l.startswith("| ") and not l.startswith("| Bucket")}


@pytest.mark.parametrize("journal_id", ["nature", "nature-medicine", "cancer-cell", "ieee-tmi", "ieee-jbhi"])
@pytest.mark.parametrize("free_text", [False, True])
def test_report_body_row_for_real_entries(tmp_path, journal_id, free_text):
    d = synth_manuscript(tmp_path)
    extra = []
    jid = journal_id
    if free_text:
        extra = ["--overrides", str(write_overrides(tmp_path, free_text_entry(journal_id)))]
        jid = f"{journal_id}-ft"
    r = run_export(d, jid, tmp_path / "x.docx", "--report-only", *extra)
    assert r.returncode == 0, r.stderr
    count, limit, status = report_rows(r.stdout)["Body"]
    if journal_id.startswith("ieee"):
        assert limit == "—" and "| Pages | — | ≤" in r.stdout and "超出" not in status
    else:
        assert count == "3000", r.stdout          # Methods (3000 words) excluded, as the limits say
        assert "超出" not in status and limit != "≤6"


def test_free_text_reading():
    m = _load_module()
    rd = m._read_limit
    assert rd("more than 30 expected for research articles")["value"] is None
    assert "minimum" in rd("more than 30 expected for research articles")["note"]
    assert rd("no fixed maximum (each counts 250 words against the limit)", "count")["no_limit"]
    assert rd("no maximum (text must be concise and comprehensive)")["no_limit"]
    assert rd("16 pages")["pages"] == 16 and rd("16 pages")["value"] is None
    assert rd("8–10 printed pages including figures")["pages"] == 10
    ieee = ("≤10 pages (double column) at initial submission — longer papers returned without review; "
            "overlength charges beyond 8 published pages ($250/page for 9-10, $350/page 11+)")
    assert rd(ieee)["pages"] == 10 and rd(ieee)["value"] is None
    lancet = rd("3000 (randomised trials may be extended to 4500; observational studies and systematic "
                "reviews/meta-analyses: 3500)")
    assert lancet["value"] == 3000 and "4500" in lancet["note"]
    assert rd("≤60 (not strictly enforced)", "count") == {"value": 60, "soft": True, "pages": None,
                                                         "no_limit": False, "note": None}
    assert rd("≤40 (counted within the 4000-word limit), Vancouver", "count")["note"] is None
    assert rd("no fixed limit (≥600 dpi PNG/JPEG/TIFF; color free of charge)", "count")["no_limit"]
    rng = rd("5000–6000 (Basic Science / Clinical Research Articles)")   # which bound for which type?
    assert rng["value"] is None and "range" in rng["note"]
    assert rd("150-250 words")["value"] == 250                           # a plain min–max range
    parts = rd("≤20 (figure pieces and/or tables combined)", "count")
    assert parts["value"] is None and "figure parts" in parts["note"]
    assert rd("≤6 (figures + tables combined, not counting figure parts)", "count")["value"] == 6


@pytest.mark.parametrize("value,limit", [
    ("3000 words (no limit on references)", 3000),
    ("4000 (excluding references, which are not limited)", 4000),
    ("≤6 (see author guidelines for details)", 6),
    ("≤40 (no limit for reviews)", 40),
    ("3500 words; no word limit for case reports", 3500),
])
def test_free_text_maximum_survives_a_later_no_limit_phrase(value, limit):
    # the "no limit" / "see author guidelines" after the number is about something else
    rd = _load_module()._read_limit(value, "count")
    assert (rd["value"], rd["no_limit"], rd["soft"]) == (limit, False, False)


def test_see_author_guidelines_alone_is_check_manually_not_no_limit():
    m = _load_module()
    rd = m._read_limit("see author guidelines", "count")
    assert rd["value"] is None and not rd["no_limit"] and "author guidelines" in rd["note"]
    assert m._read_limit("see author guidelines (no fixed limit)", "count")["no_limit"]   # urologic-oncology


@pytest.mark.parametrize("value,part,limit", [
    ("no limit for reviews; ≤40 for original articles", "references", 40),   # another article type
    ("not limited for letters, 3000 words for original articles", "words", 3000),
    ("no limit for reviews and ≤40 for original articles", "references", 40),
    ("no limit on references; 3000 words", "words", 3000),                    # another part
    ("no reference limit; 3000 words", "words", 3000),
    ("unlimited references; 3000 words", "words", 3000),
    ("no word limit; ≤40", "references", 40),                               # "words": the word limit
    ("references not limited; ≤6 figures", "figures", 6),
    ("no limit on supplementary figures; ≤6", "figures", 6),
])
def test_no_limit_phrase_about_another_type_or_part_does_not_cancel_the_limit(value, part, limit):
    rd = _load_module()._read_limit(value, "count", part)
    assert (rd["value"], rd["no_limit"]) == (limit, False), rd


@pytest.mark.parametrize("value,part", [
    ("no limit in the guide; ≤40", "references"),                 # about what?
    ("references not limited; ≤6 figures", None),                 # part unknown
    ("≤40 for reviews; no limit for original articles", "references"),
    ("≤40 (original articles: unlimited)", "references"),
    ("no limit for reviews", "references"),                       # only another type: nothing for this one
])
def test_unclear_no_limit_phrase_is_check_manually_never_no_limit(value, part):
    rd = _load_module()._read_limit(value, "count", part)
    assert rd["value"] is None and not rd["no_limit"] and rd["note"], rd


@pytest.mark.parametrize("value,part", [
    ("references not limited; ≤6 figures", "references"), ("no reference limit for original articles", "references"),
    ("no limit for reviews and original articles", "references"), ("no word limit", "abstract"),
    ("no limit on figures or tables", "tables"), ("unlimited in number", "figures"),
    ("unlimited within page limit", "figures"), ("no strict limit (reviews)", "words"),
    ("no strict limit stated in guide for authors (keep concise)", "words"),
])
def test_no_limit_phrase_about_this_part_is_no_limit(value, part):
    assert _load_module()._read_limit(value, "count", part)["no_limit"]


def test_no_limit_for_another_type_or_part_in_overrides_free_text_report(tmp_path):
    ov = write_overrides(tmp_path, {"id": "scoped", "journal": "Scoped", "guidelines_checked": "2026-09",
                                    "word_limit": "no limit on references; 3000 words",
                                    "references": "no limit for reviews; ≤40 for original articles",
                                    "figures": "no limit in the guide for authors", "tables": "≤3"})
    refs = "# References\n\n" + "".join(f"{i}. Ref {i}.\n" for i in range(1, 46))
    d = synth_manuscript(tmp_path, {"results.md": "# Results\n\n" + " ".join(["w"] * 3100) + "\n",
                                    "references.md": refs})
    r = run_export(d, "scoped", tmp_path / "x.docx", "--report-only", "--overrides", str(ov))
    assert r.returncode == 0, r.stderr
    rows = report_rows(r.stdout)
    assert rows["Body"] == ["3100", "≤3000", "⚠️ 超出 100"], r.stdout
    assert rows["References"] == ["45", "≤40", "⚠️ 超出 5"], r.stdout
    assert rows["Figures"][1:] == ["—", "ℹ️ no limit in the library — check the journal's guidelines"], r.stdout
    assert "no fixed limit (template)" not in r.stdout
    assert "Check manually — `figures`: the template says 'no limit in the guide for authors'" in r.stdout


def test_body_parts_named_in_free_text():
    m = _load_module()
    ex, inc = m._body_parts_from_text("7000 (including main figure legends; excluding STAR Methods, "
                                      "supplemental legends and references)")
    assert ex == {"methods", "references"} and inc == {"figure_legends"}
    ex, inc = m._body_parts_from_text("4000 (excluding abstract, Methods, references and figure legends)")
    assert ex == {"abstract", "methods", "references", "figure_legends"} and not inc
    assert m._body_parts_from_text("3000 (including abstract)") == (set(), {"abstract"})
    assert m._body_parts_from_text("6000", "no fixed limit (counted in 6000-word limit), AMA") == \
        (set(), {"references"})


def test_lancet_family_combined_display_items_from_free_text(mod, tmp_path):
    # lancet-oncology: figures "≤5", tables "included in figure limit" — tables count against the 5
    ov = write_overrides(tmp_path, free_text_entry("lancet-oncology"))
    md = ("# Results\n\n**Figure 1.** A.\n\n**Figure 2.** B.\n\n**Figure 3.** C.\n\n"
          + "".join(f"| t{i} | x |\n|---|---|\n| 1 | 2 |\n\n" for i in range(3)))
    d = synth_manuscript(tmp_path, {"results.md": md})
    r = run_export(d, "lancet-oncology-ft", tmp_path / "x.docx", "--report-only", "--overrides", str(ov))
    assert r.returncode == 0, r.stderr
    row = next(l for l in r.stdout.splitlines() if l.startswith("| Display items"))
    assert "| 6 | ≤5 | ⚠️ 超出 1 |" in row


PER_ITEM_TEXT = "3000 (excluding abstract and references; each table or figure counts as 250 words)"
PER_ITEM_MS = {"results.md": "# Results\n\n" + " ".join(["w"] * 2900) + "\n\n"
               + "".join(f"**Table {i}.** T\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n" for i in range(1, 4))}


def test_each_table_or_figure_counts_n_words(tmp_path):
    # world-journal-urology rule: 2900 words of text + 3 tables × 250 = 3650 > 3000
    ov = write_overrides(tmp_path, {"id": "per-item", "journal": "Per Item", "guidelines_checked": "2026-09",
                                    "word_limit": PER_ITEM_TEXT, "limits": {"body_words": 3000}})
    d = synth_manuscript(tmp_path, PER_ITEM_MS)
    r = run_export(d, "per-item", tmp_path / "x.docx", "--report-only", "--overrides", str(ov))
    rows = report_rows(r.stdout)
    assert rows["Body"] == ["2900", "≤3000", "ℹ️ text only — see the next row"]       # never a plain ✅
    assert rows["Body + 250 words per table / figure"] == ["3650", "≤3000", "⚠️ 超出 650"], r.stdout
    assert "(template rule; 3 tables / figures)" in r.stdout


def test_each_table_or_figure_rule_is_not_applied_to_other_article_types(tmp_path):
    ov = write_overrides(tmp_path, {"id": "per-item", "journal": "Per Item", "guidelines_checked": "2026-09",
                                    "word_limit": PER_ITEM_TEXT, "limits": {"body_words": 3000},
                                    "limits_by_type": {"review": {"body_words": 5000}}})
    d = synth_manuscript(tmp_path, PER_ITEM_MS)
    r = run_export(d, "per-item", tmp_path / "x.docx", "--report-only", "--overrides", str(ov),
                   "--article-type", "review")
    assert report_rows(r.stdout)["Body"] == ["2900", "≤5000", "✅"]
    assert "words per table / figure" not in r.stdout


@pytest.mark.parametrize("free_text", [False, True])
def test_each_table_or_figure_rule_with_the_real_entry(tmp_path, free_text):
    # world-journal-urology: checked against 3000 when a body_words is known (free text, or a library
    # `limits.body_words`); without one the total is shown with "no limit in the library", never ✅
    extra, jid = [], "world-journal-urology"
    if free_text:
        extra, jid = ["--overrides", str(write_overrides(tmp_path, free_text_entry(jid)))], f"{jid}-ft"
    r = run_export(synth_manuscript(tmp_path, PER_ITEM_MS), jid, tmp_path / "x.docx", "--report-only", *extra)
    assert r.returncode == 0, r.stderr
    rows = report_rows(r.stdout)
    total, row = int(rows["Body"][0]) + 3 * 250, rows["Body + 250 words per table / figure"]
    limit = 3000 if free_text else (LIBRARY["world-journal-urology"].get("limits") or {}).get("body_words")
    if limit is not None:
        assert row == [str(total), f"≤{limit}", f"⚠️ 超出 {total - limit}"], r.stdout
    else:
        assert row[0] == str(total) and row[1] == "—" and "no limit in the library" in row[2]
        assert "the library entry has no `body_words` — compare the 'Body + 250 words per table" in r.stdout
    assert "✅" not in rows["Body"][2]


def test_malformed_limits_block_exits_2(manuscript, tmp_path):
    ov = write_overrides(tmp_path, {"id": "mine", "journal": "Mine", "limits": {"body_words": "3000 words"}})
    r = run_export(manuscript, "mine", tmp_path / "x.docx", "--overrides", str(ov))
    assert r.returncode == 2 and "Traceback" not in r.stderr
    assert r.stderr.strip().startswith("Error: overrides file, entry `mine`: `limits.body_words`")


def test_structured_limits_win_over_free_text(mod, tmp_path):
    ov = write_overrides(tmp_path, {"id": "mine", "journal": "Mine", "word_limit": "no strict limit",
                                    "limits": {"body_words": 2500, "body_excludes": ["abstract", "references"],
                                               "references": 10}})
    config = mod.get_journal_config("mine", str(YAML), str(ov))
    assert config["limits"]["body_words"] == 2500 and config["limits_source"] == "journal-overrides `limits`"
    assert config["body_rule"] == "library"


def test_soft_limits_are_information_not_warnings(tmp_path):
    ov = write_overrides(tmp_path, {"id": "soft-j", "journal": "Soft", "guidelines_checked": "2026-09",
                                    "limits": {"body_words": 1000, "soft": True}})
    d = synth_manuscript(tmp_path, {"results.md": "# Results\n\n" + " ".join(["w"] * 1200) + "\n"})
    r = run_export(d, "soft-j", tmp_path / "x.docx", "--report-only", "--overrides", str(ov))
    assert report_rows(r.stdout)["Body"] == ["1200", "≤1000", "ℹ️ 200 above the guidance value (not a hard limit)"]


def test_body_counts_what_body_excludes_does_not_exclude(tmp_path):
    # figure legends and the abstract count when body_excludes leaves them out
    ov = write_overrides(tmp_path, {"id": "incl", "journal": "Incl", "guidelines_checked": "2026-09",
                                    "limits": {"body_words": 5000, "body_excludes": ["references", "tables"]}})
    d = synth_manuscript(tmp_path, {"abstract.md": "# Abstract\n\none two three\n",
                                    "results.md": "# Results\n\nfour five\n\n**Figure 1.** six seven eight\n"})
    r = run_export(d, "incl", tmp_path / "x.docx", "--report-only", "--overrides", str(ov))
    assert report_rows(r.stdout)["Body"][0] == "10"           # 2 + abstract 3 + legend "Figure 1. six seven eight" 5
    assert "Body (results + abstract + figure legends)" in r.stdout


# ─── --article-type and unverified entries (D1 / D2) ───────────────────────

def test_article_type_without_library_limits_is_not_checked(tmp_path):
    d = synth_manuscript(tmp_path)
    r = run_export(d, "bju-international", tmp_path / "x.docx", "--report-only", "--article-type", "review")
    assert r.returncode == 0, r.stderr
    assert "The library has no review limits for BJU International" in r.stdout
    counts = r.stdout.split("## Counts")[1].split("## Sections")[0]
    assert "✅" not in counts and "超出" not in counts and "ℹ️ not checked" in counts


def test_article_type_uses_limits_by_type(tmp_path):
    ov = write_overrides(tmp_path, {"id": "typed", "journal": "Typed", "guidelines_checked": "2026-09",
                                    "limits": {"body_words": 3000},
                                    "limits_by_type": {"research_letter": {"body_words": 600}}})
    d = synth_manuscript(tmp_path, {"results.md": "# Results\n\n" + " ".join(["w"] * 800) + "\n"})
    r = run_export(d, "typed", tmp_path / "x.docx", "--report-only", "--overrides", str(ov),
                   "--article-type", "research-letter")
    assert report_rows(r.stdout)["Body"] == ["800", "≤600", "⚠️ 超出 200"]
    assert "limits_by_type.research_letter" in r.stdout


def test_no_fixed_limit_label_comes_from_the_original_article_text_only(tmp_path):
    ov = write_overrides(tmp_path, {"id": "nl", "journal": "NL", "guidelines_checked": "2026-09",
                                    "references": "unlimited", "limits": {"body_words": 3000},
                                    "limits_by_type": {"review": {"body_words": 6000}}})
    d = synth_manuscript(tmp_path, {"results.md": "# Results\n\nText.\n", "references.md": "1. A.\n"})
    r = run_export(d, "nl", tmp_path / "x.docx", "--report-only", "--overrides", str(ov))
    assert report_rows(r.stdout)["References"] == ["1", "—", "ℹ️ no fixed limit (template)"]
    r = run_export(d, "nl", tmp_path / "x.docx", "--report-only", "--overrides", str(ov), "--article-type", "review")
    assert report_rows(r.stdout)["References"] == ["1", "—", "ℹ️ no limit in the library — check the journal's guidelines"]


_NO_LIMIT_ROWS = [("word_limit", "no strict limit", "Body"), ("abstract", "structured, no word limit", "Abstract"),
                  ("references", "unlimited, Vancouver", "References"), ("figures", "unlimited", "Figures"),
                  ("tables", "no fixed limit", "Tables")]


@pytest.mark.parametrize("field,value,row,limits_block",
                         [(*r, b) for r in _NO_LIMIT_ROWS for b in (False, True)]
                         + [("figures", "no fixed limit (figures and tables combined)", "Display items", False)])
def test_no_fixed_limit_label_on_each_row(tmp_path, field, value, row, limits_block):
    # free text only, or a `limits` block without this key: the label comes from the template text
    entry = {"id": "nl", "journal": "NL", "guidelines_checked": "2026-09", field: value}
    if row == "Display items":
        entry["tables"] = "included in figure limit"
    if limits_block:
        entry["limits"] = {"references": 99} if row != "References" else {"body_words": 99999}
    d = synth_manuscript(tmp_path, {"abstract.md": "# Abstract\n\nShort.\n", "results.md": "# Results\n\nText.\n",
                                    "references.md": "1. A.\n"})
    r = run_export(d, "nl", tmp_path / "x.docx", "--report-only", "--overrides", str(write_overrides(tmp_path, entry)))
    assert r.returncode == 0, r.stderr
    rows = report_rows(r.stdout)
    assert rows[row][1:] == ["—", "ℹ️ no fixed limit (template)"], r.stdout
    assert sum("no fixed limit (template)" in line for line in r.stdout.splitlines()) == 1, r.stdout


@pytest.mark.parametrize("journal_id,article_type", [("bmj", "systematic_review"),
                                                     ("european-journal-radiology", "brief_report")])
def test_real_entries_other_article_types_never_say_no_fixed_limit(tmp_path, journal_id, article_type):
    # bmj "no fixed limit (research articles …)" and EJR "see author guidelines (no fixed limit)"
    # describe the original article, not this type
    d = synth_manuscript(tmp_path, {"results.md": "# Results\n\nText.\n"})
    r = run_export(d, journal_id, tmp_path / "x.docx", "--report-only", "--article-type", article_type)
    assert r.returncode == 0, r.stderr
    counts = r.stdout.split("## Counts")[1].split("## Sections")[0]
    assert "no fixed limit (template)" not in counts


def test_bad_article_type_is_one_line_exit_2(manuscript, tmp_path):
    r = run_export(manuscript, "nature", tmp_path / "x.docx", "--article-type", "editorial")
    assert r.returncode == 2 and "Traceback" not in r.stderr and len(r.stderr.strip().splitlines()) == 1


def test_unverified_entry_warns_and_marks_passing_checks(tmp_path):
    entry = {"id": "unver", "journal": "Unverified J", "limits": {"body_words": 5000, "references": 50}}
    d = synth_manuscript(tmp_path, {"results.md": "# Results\n\nshort text\n"})
    r = run_export(d, "unver", tmp_path / "x.docx", "--report-only", "--overrides",
                   str(write_overrides(tmp_path, entry)))
    counts = r.stdout.split("## Counts")[1]
    assert counts.split("\n")[1].startswith("- ⚠️ **Unverified journal entry:** limits not checked against "
                                            "the journal's current author guidelines")
    assert report_rows(r.stdout)["Body"] == ["2", "≤5000", "✅*"]
    assert "✅* = within the library value" in r.stdout
    before_report = r.stdout.split("# Manuscript Export Report")[0].splitlines()   # the console line (D2)
    assert ("⚠️ Unverified J: limits not checked against the journal's current author guidelines; "
            "confirm on the journal site before relying on them") in before_report
    entry["guidelines_checked"] = "2026-09"
    r = run_export(d, "unver", tmp_path / "x.docx", "--report-only", "--overrides",
                   str(write_overrides(tmp_path, entry)))
    assert report_rows(r.stdout)["Body"] == ["2", "≤5000", "✅"] and "Unverified journal entry" not in r.stdout
    assert "not checked against the journal" not in r.stdout


def test_pending_items_are_named_in_the_limit_section(tmp_path):
    # JD-4: a 待核实 value is null in `limits` (its row says "no limit in the library"), so the limit
    # section itself names it; before, it was only in the special-requirements list at the end.
    entry = {"id": "pend", "journal": "Pending J", "guidelines_checked": "2026-09",
             "limits": {"body_words": 5000, "references": None},
             "special": ["待核实: reference limit (≤60) not stated in the guidelines", "Cover letter required"]}
    d = synth_manuscript(tmp_path, {"results.md": "# Results\n\nshort text\n"})
    r = run_export(d, "pend", tmp_path / "x.docx", "--report-only", "--overrides",
                   str(write_overrides(tmp_path, entry)))
    assert r.returncode == 0, r.stderr
    counts = r.stdout.split("## Counts")[1].split("## Sections")[0]
    line = next(ln for ln in counts.splitlines() if "Unconfirmed in the library" in ln)
    assert line == ("- ⚠️ **Unconfirmed in the library (待核实):** reference limit (≤60) not stated in the "
                    "guidelines — confirm on the journal site.")
    entry["special"] = ["Cover letter required"]
    r = run_export(d, "pend", tmp_path / "x.docx", "--report-only", "--overrides",
                   str(write_overrides(tmp_path, entry)))
    assert r.returncode == 0 and "Unconfirmed in the library" not in r.stdout
    r = run_export(d, "nature-medicine", tmp_path / "x.docx", "--report-only")       # a real checked entry
    assert r.returncode == 0 and "Unconfirmed in the library (待核实):** abstract length" in r.stdout


# ─── journal-neutral export (SB-9) ─────────────────────────────────────────

def test_generic_journal_exports_without_limit_checks(manuscript, tmp_path):
    out = tmp_path / "generic.docx"
    r = run_export(manuscript, "generic", out, "--report", str(tmp_path / "g.md"))
    assert r.returncode == 0, r.stderr
    assert out.exists()
    h = headings(out)
    assert h.index("Introduction") < h.index("Methods") < h.index("Results") < h.index("Discussion")
    counts = r.stdout.split("## Counts")[1].split("## Sections")[0]
    assert "Journal-neutral export" in counts and "✅" not in counts and "超出" not in counts
    assert "⚠️ **Missing:**" not in r.stdout             # key points / research in context are optional


# ─── figure / table / reference counting (OS-2, OS-9, CG-2) ─────────────────

LEGEND_STYLES = ("**Figure 1. Kaplan–Meier curves for OS.**\n"
                 "Figure 2. Plain style.\n"
                 "**Figure 3a.** Panel a\n\n"
                 "## Figure 4. Heading style\n\n"
                 "Fig. 5 | Nature style\n\n"
                 "图6 流程图\n\n"
                 "Figure 1 shows running text, not a legend.\n")


def test_common_legend_styles_are_counted(mod, tmp_path):
    _, stats, report = export_md(mod, tmp_path, LEGEND_STYLES, "figure-legends.md", journal="bju-international")
    assert stats["fig_numbers"] == {1, 2, 3, 4, 5, 6}
    row = next(l for l in report.splitlines() if l.startswith("| Figures"))
    assert "| 6 | ≤3 | ⚠️ 超出 3 |" in row


@pytest.mark.parametrize("style", ["[{i}] Author A. Title {i}. J. 2020.", "({i}) Author A. Title {i}. J. 2020.",
                                   "{i}. Author A. Title {i}. J. 2020.", "- Author A. Title {i}. J. 2020."])
def test_reference_styles_are_counted(mod, tmp_path, style):
    md = "# References\n\n" + "\n".join(style.format(i=i) for i in range(1, 61)) + "\n"
    doc, stats, report = export_md(mod, tmp_path, md, "references.md", journal="bju-international")
    assert stats["ref_count"] == 60
    assert "| References (entries) | 60 | ≤30 | ⚠️ 超出 30 |" in report
    assert len([p for p in doc.paragraphs if "Author A." in p.text]) == 60     # one paragraph per reference


def test_wrapped_reference_line_starting_with_a_year_is_not_a_new_entry(mod, tmp_path):
    md = "# References\n\n[1] Smith J. A long title that wraps\n(2020) onto a line starting with a year.\n[2] Doe A.\n"
    _, stats, _ = export_md(mod, tmp_path, md, "references.md")
    assert stats["ref_count"] == 2 and stats["ref_unnumbered"] == 0


def test_unnumbered_references_are_never_a_plain_pass(mod, tmp_path):
    md = "# References\n\nSmith J (2020) Title. J.\n\nDoe A (2021) Title. J.\n"
    _, _, report = export_md(mod, tmp_path, md, "references.md", journal="bju-international")
    row = next(l for l in report.splitlines() if l.startswith("| References"))
    assert "❓" in row and "✅" not in row


def test_unrecognised_legends_file_is_not_a_pass(mod, tmp_path):
    _, _, report = export_md(mod, tmp_path, "Kaplan–Meier curves.\n\nForest plot.\n", "figure-legends.md",
                             journal="bju-international")
    row = next(l for l in report.splitlines() if l.startswith("| Figures"))
    assert "❓" in row and "✅" not in row


def test_cited_figure_numbers_count(mod, tmp_path):
    md = ("# Results\n\nSee Figures 2–4 and Supplementary Figure 9, eFigure 8, Figure S7; 如图5所示。"
          "Data from Figure 2018 onwards.\n")
    _, stats, report = export_md(mod, tmp_path, md, journal="bju-international")
    assert stats["fig_cited"] == {2, 4, 5}
    assert "| 5 | ≤3 | ⚠️ 超出 2 |" in next(l for l in report.splitlines() if l.startswith("| Figures"))


@pytest.mark.parametrize("md,count", [
    ("**Table 1.** A\n\n**Table 2.** B\n\n**Table 3.** C\n\n**Table 4.** D\n", 4),   # captions, no Markdown table
    ("# Results\n\nAs shown in Table 4, rates fell.\n", 4),                          # the highest number cited
])
def test_table_count_uses_captions_and_citations(mod, tmp_path, md, count):
    _, stats, report = export_md(mod, tmp_path, md, journal="bju-international")
    assert stats["table_count"] == 0
    row = next(l for l in report.splitlines() if l.startswith("| Tables"))
    assert f"| {count} | ≤3 | ⚠️ 超出 {count - 3} |" in row


def test_table_captions_in_the_legends_file_count(mod, tmp_path):
    # captions in body files are also citations; in figure-legends.md only the caption term counts them
    md = "**Figure 1.** Flow.\n\n" + "".join(f"**Table {i}.** Title {i}\n\n" for i in range(1, 5))
    _, stats, report = export_md(mod, tmp_path, md, "figure-legends.md", journal="bju-international")
    assert stats["table_numbers"] == {1, 2, 3, 4} and not stats["table_cited"] and stats["table_count"] == 0
    row = next(l for l in report.splitlines() if l.startswith("| Tables"))
    assert "captions: 4" in row and "| 4 | ≤3 | ⚠️ 超出 1 |" in row


@pytest.mark.parametrize("text", ["既往已发表5篇相关研究。", "红色代表4级", "研究者试图9次", "列表3项", "附图2", "补充表3",
                                  "发表 5篇", "代表 4 级", "地图3张", "意图2个", "企图5次", "蓝图2", "版图3", "示意图2", "截图2张",
                                  "图表3", "时间表2", "时刻表4", "统计表3份", "量表3个维度", "外表5", "补充材料表1", "扩展数据图1"])
def test_cjk_words_ending_in_tu_or_biao_are_not_citations(text):
    m = _load_module()
    assert not m.cited_numbers(text, m._FIG_CITE_RE) and not m.cited_numbers(text, m._TAB_CITE_RE)


@pytest.mark.parametrize("text,fig,tab", [("如表2所示", set(), {2}), ("结果见图3。", {3}, set()), ("表1和表2", set(), {1, 2}),
                                          ("（图4）", {4}, set()), ("与图5比较", {5}, set()), ("。表6显示", set(), {6}),
                                          ("如图1~3所示", {1, 3}, set()), ("如下表2", set(), {2})])
def test_cjk_figure_and_table_citations_are_counted(text, fig, tab):
    m = _load_module()
    assert (m.cited_numbers(text, m._FIG_CITE_RE), m.cited_numbers(text, m._TAB_CITE_RE)) == (fig, tab)


@pytest.mark.parametrize("text,fig,tab", [("对图3进行分析", {3}, set()), ("结合表2可见", set(), {2}),
                                          ("通过图4可以看出", {4}, set()), ("将表2中的数据", set(), {2}),
                                          ("为图1所示", {1}, set()), ("观察图4", {4}, set()), ("比较图2与图3", {2, 3}, set()),
                                          ("此外表5显示", set(), {5}), ("注意图3", {3}, set()), ("样本量表2", set(), {2})])
def test_cjk_citations_after_any_word_are_counted(text, fig, tab):
    # 图N / 表N is recognised by itself, not by a list of the words that may come before it
    m = _load_module()
    assert (m.cited_numbers(text, m._FIG_CITE_RE), m.cited_numbers(text, m._TAB_CITE_RE)) == (fig, tab)


def test_chinese_results_citing_one_more_figure_and_table_are_over_the_limit(tmp_path):
    d = synth_manuscript(tmp_path, {
        "figure-legends.md": "**Figure 1.** Flow.\n\n**Figure 2.** ROC.\n\n**Figure 3.** Calibration.\n",
        "results.md": ("# 结果\n\n如图1所示，共纳入120例患者。图2显示了ROC曲线。见图3。通过图4可以看出，模型表现稳定。"
                       "对表4中的亚组进行了分析。\n\n"
                       + "".join(f"**Table {i}.** T\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n" for i in range(1, 4)))})
    r = run_export(d, "bju-international", tmp_path / "x.docx", "--report-only")
    assert r.returncode == 0, r.stderr
    rows = report_rows(r.stdout)
    assert rows["Figures"] == ["4", "≤3", "⚠️ 超出 1"] and rows["Tables"] == ["4", "≤3", "⚠️ 超出 1"], r.stdout
    assert "The text cites Figure 4, but legends were found only for Figure 1, 2, 3" in r.stdout


def test_chinese_introduction_without_tables_is_not_over_the_table_limit(tmp_path):
    d = synth_manuscript(tmp_path, {"introduction.md": "# 引言\n\n既往已发表5篇相关研究，红色代表4级，研究者试图9次。\n"})
    r = run_export(d, "bju-international", tmp_path / "x.docx", "--report-only")
    rows = report_rows(r.stdout)
    assert rows["Tables"][0] == "0" and rows["Figures"][0] == "0" and "超出" not in r.stdout


def test_missing_abstract_is_not_a_pass(mod, tmp_path):
    _, _, report = export_md(mod, tmp_path, "# Results\n\nText.\n", journal="bju-international")
    row = next(l for l in report.splitlines() if l.startswith("| Abstract"))
    assert "❓" in row and "abstract.md not found" in row and "✅" not in row


def test_images_in_comments_and_code_are_not_counted(mod, tmp_path):
    md = "# Results\n\n<!-- ![old](a.png) -->\n\n```\n![code](b.png)\n```\n\nText `![x](c.png)` and ![Fig](d.png).\n"
    _, stats, _ = export_md(mod, tmp_path, md)
    assert stats["fig_embedded"] == 1


def test_legends_in_project_folder_are_read(manuscript, tmp_path):
    (tmp_path / "figure-legends.md").write_text("**Figure 1.** One.\n\n**Figure 2.** Two.\n", encoding="utf-8")
    out = tmp_path / "x.docx"
    r = run_export(manuscript, "european-urology", out)          # cwd = project folder (tmp_path)
    assert r.returncode == 0, r.stderr
    assert "Two." in full_text(out)
    assert "the legends were read from `figure-legends.md` (the project folder" in r.stdout
    (manuscript / "figure-legends.md").write_text("**Figure 1.** Inside.\n", encoding="utf-8")
    r = run_export(manuscript, "european-urology", out)
    assert "Inside." in full_text(out) and "Two." not in full_text(out)
    assert "also exists and was NOT exported" in r.stdout


def test_legends_of_another_project_in_the_current_folder_are_not_read(tmp_path):
    d = synth_manuscript(tmp_path / "X", {"results.md": "# Results\n\nText.\n"})
    other = tmp_path / "X" / "otherproj"
    other.mkdir()
    (other / "figure-legends.md").write_text("**Figure 1.** OTHER PROJECT legend\n", encoding="utf-8")
    out = tmp_path / "x.docx"
    r = run_export(d, "european-urology", out, cwd=other)
    assert r.returncode == 0, r.stderr
    assert "OTHER PROJECT" not in full_text(out) and "legends were read from" not in r.stdout
    assert "Optional sections not present: conclusion, figure-legends" in r.stdout


def test_legends_in_the_current_folder_that_contains_the_manuscript_are_read(tmp_path):
    proj = tmp_path / "proj"
    d = synth_manuscript(proj / "paper", {"results.md": "# Results\n\nText.\n"})       # proj/paper/manuscript
    (proj / "figure-legends.md").write_text("**Figure 1.** Project legend.\n", encoding="utf-8")
    out = tmp_path / "x.docx"
    r = run_export(d, "european-urology", out, cwd=proj)
    assert r.returncode == 0, r.stderr
    assert "Project legend." in full_text(out)
    assert "(the current folder, which contains the manuscript folder)" in r.stdout


# ─── words (OS-8) ──────────────────────────────────────────────────────────

def test_cjk_characters_are_counted_one_by_one(mod, tmp_path):
    _, stats, report = export_md(mod, tmp_path, "# 结果\n\n本研究纳入了一百名患者。AUC 为 0.90。\n")
    assert stats["words"]["results"] == 12 + 2              # 12 Han characters + "AUC" + "0.90"
    assert "each Chinese / Japanese character counts as one word" in report


def test_legends_and_captions_in_body_files_are_not_body_words(mod, tmp_path):
    md = ("# Results\n\nOne two three.\n\n**Figure 1.** " + " ".join(["legend"] * 200)
          + "\n\n**Table 1.** Baseline characteristics\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    _, stats, report = export_md(mod, tmp_path, md)
    assert stats["words"]["results"] == 3 and stats["legend_words"] == 202     # "Figure 1." + 200
    assert stats["fig_numbers"] == {1} and stats["table_numbers"] == {1}
    assert "| Figure legends inside body files | 202 |" in report


# ─── tildes (OS-4) ─────────────────────────────────────────────────────────

@pytest.mark.parametrize("text", ["年龄18~65岁，体重50~80kg。", "range 1~2,3~4", "IQR 2~5;7~9",
                                  "中位随访36~48个月，范围12~60个月", "TNM分期I~II期和III~IV期", "grade I~II,III~IV",
                                  "dose 5~10mg/10~20mg", "ratio 1~2/3~4"])
def test_tilde_ranges_are_not_subscripts(text):
    m = _load_module()
    assert m.plain_text(text) == text
    assert not any(a.get("sub") for _, a in m.parse_inline(text))


def test_real_subscripts_still_work_and_are_reported(mod, tmp_path):
    assert mod.parse_inline("CO~2~ and T~1~-weighted") == [("CO", {}), ("2", {"sub": True}), (" and T", {}),
                                                         ("1", {"sub": True}), ("-weighted", {})]
    _, _, report = export_md(mod, tmp_path, "# Methods\n\nH~2~O, age 18~65.\n", "methods.md")
    assert "methods.md:3 — `~x~` set as subscript" in report and "`~2~`" in report


# ─── BOM, placeholders, unsupported Markdown (OS-3, OS-6, OS-7) ─────────────

def test_utf8_bom_is_ignored(mod, tmp_path):
    d = tmp_path / "ms"
    d.mkdir()
    (d / "introduction.md").write_text("# Introduction\n\nText here.\n", encoding="utf-8-sig")
    (d / "results.md").write_text("| a | b |\n|---|---|\n| 1 | 2 |\n", encoding="utf-8-sig")
    config = mod.get_journal_config("european-urology", str(YAML))
    out = tmp_path / "bom.docx"
    mod.build_docx(str(d), config, str(out))
    doc = Document(str(out))
    assert ("Heading 1", "Introduction") in paras(doc)
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["a", "b"], ["1", "2"]]


@pytest.mark.parametrize("snippet", ["<!--\nTODO: add the AUC\n-->", "AUC was TBD in the test set.",
                                     "AUC（待补充）。", "Sites: TODO", "随访时间待核实。", "as shown [ref].",
                                     "FIXME: wording"])
def test_more_placeholders_are_reported(mod, tmp_path, snippet):
    _, stats, report = export_md(mod, tmp_path, f"# Results\n\n{snippet}\n")
    assert stats["placeholders"] and "No placeholders detected" not in report


def test_blinding_xxx_and_lowercase_words_are_not_placeholders(mod, tmp_path):
    _, stats, _ = export_md(mod, tmp_path, "# Results\n\nPatients from Hospital XXX; a todo list; tbd.\n")
    assert not stats["placeholders"]


def test_removed_comments_are_noted(mod, tmp_path):
    _, _, report = export_md(mod, tmp_path, "# Results\n\nText <!-- reviewer 2 asked -->.\n")
    assert "results.md:3 — HTML comment removed (not exported)" in report


def test_unsupported_markdown_is_reported(mod, tmp_path):
    md = ("# Results\n\nText[^1] see [the site][ref] and <span>x</span>, $\\alpha = 0.05$.\n\n"
          "[^1]: Footnote.\n\n[ref]: https://x.org\n\n- ## heading in list\n\n"
          "It costs $250 and $350 per page.\n")
    _, _, report = export_md(mod, tmp_path, md)
    notes = report.split("## Conversion notes")[1].split("\n## ")[0]
    for msg in ("Markdown footnote", "reference-style link definition", "reference-style link `[text][id]`",
                "inline HTML", "TeX math", "heading inside a list item"):
        assert msg in notes, msg
    assert "results.md:11" not in notes                  # "$250 and $350" is not math


# ─── review section files (OS-10) ──────────────────────────────────────────

def test_review_section_files_are_exported_in_order(tmp_path):
    d = synth_manuscript(tmp_path, {"introduction.md": "# Introduction\n\nIntro.\n",
                                    "section-10-outlook.md": "# Outlook\n\nTen.\n",
                                    "section-2-ai.md": "# AI\n\nTwo two.\n",
                                    "section-1-imaging.md": "# Imaging\n\nOne one one.\n",
                                    "conclusion.md": "# Conclusion\n\nEnd.\n"})
    out = tmp_path / "rev.docx"
    r = run_export(d, "bju-international", out)
    assert r.returncode == 0, r.stderr
    assert headings(out) == ["Introduction", "Imaging", "AI", "Outlook", "Conclusion"]
    assert report_rows(r.stdout)["Body"][0] == "8"
    assert "not exported" not in r.stdout and "**Missing:** `methods.md`" not in r.stdout


# ─── page and line numbers (CG-7) ──────────────────────────────────────────

def _docx_xml(path):
    import zipfile
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode(), "".join(
            z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))


def test_page_and_line_numbers_by_default(manuscript, tmp_path):
    out = tmp_path / "n.docx"
    r = run_export(manuscript, "european-urology", out)
    assert r.returncode == 0, r.stderr
    body, footer = _docx_xml(out)
    assert 'w:lnNumType w:countBy="1" w:restart="continuous"' in body and "PAGE" in footer
    assert "page numbers bottom centre, line numbers (continuous) — page numbers: library defaults" in r.stdout


def test_page_and_line_numbers_can_be_turned_off(manuscript, tmp_path):
    out = tmp_path / "n.docx"
    assert run_export(manuscript, "european-urology", out, "--no-line-numbers", "--no-page-numbers").returncode == 0
    body, footer = _docx_xml(out)
    assert "lnNumType" not in body and "PAGE" not in footer
    ov = write_overrides(tmp_path, {"id": "noln", "journal": "No LN", "line_numbers": False})
    assert run_export(manuscript, "noln", out, "--overrides", str(ov)).returncode == 0
    body, footer = _docx_xml(out)
    assert "lnNumType" not in body and "PAGE" in footer


# ─── output paths never overwrite inputs (SEC-12) and CLI errors (CLI-1) ─────

@pytest.mark.parametrize("flag,target", [("--report", "results.md"), ("--supplementary", "supplementary.md"),
                                         ("--output", "results.md")])
def test_outputs_never_overwrite_inputs(manuscript, tmp_path, flag, target):
    before = (manuscript / target).read_bytes()
    args = [flag, str(manuscript / target)]
    out = tmp_path / "x.docx" if flag != "--output" else manuscript / target
    r = run_export(manuscript, "european-urology", out, *(args if flag != "--output" else []))
    assert r.returncode == 2 and "Traceback" not in r.stderr and len(r.stderr.strip().splitlines()) == 1
    assert (manuscript / target).read_bytes() == before


def test_report_does_not_overwrite_other_files_but_replaces_an_earlier_report(manuscript, tmp_path):
    notes = tmp_path / "notes.md"
    notes.write_text("my notes\n", encoding="utf-8")
    r = run_export(manuscript, "european-urology", tmp_path / "x.docx", "--report", str(notes))
    assert r.returncode == 2 and notes.read_text(encoding="utf-8") == "my notes\n"
    rep = tmp_path / "export-report.md"
    assert run_export(manuscript, "european-urology", tmp_path / "x.docx", "--report", str(rep)).returncode == 0
    assert run_export(manuscript, "european-urology", tmp_path / "x.docx", "--report", str(rep)).returncode == 0


@pytest.mark.parametrize("extra", [["--report", "{out}"], ["--supplementary", "{out}"]])
def test_two_outputs_naming_the_same_file_exit_2(manuscript, tmp_path, extra):
    out = tmp_path / "same.docx"
    r = run_export(manuscript, "european-urology", out, *[a.format(out=out) for a in extra])
    assert r.returncode == 2 and "name the same file" in r.stderr and len(r.stderr.strip().splitlines()) == 1
    assert not out.exists()


@pytest.mark.parametrize("flag", ["--output", "--supplementary"])
def test_docx_outputs_must_end_in_docx(manuscript, tmp_path, flag):
    bad = tmp_path / "manuscript.md"
    args = [] if flag == "--output" else [flag, str(bad)]
    r = run_export(manuscript, "european-urology", bad if flag == "--output" else tmp_path / "x.docx", *args)
    assert r.returncode == 2 and f"{flag} must be a .docx path" in r.stderr
    assert not bad.exists() and not (tmp_path / "x.docx").exists()


def test_cli_errors_are_one_line_exit_2(manuscript, tmp_path):
    r = run_export(tmp_path / "nope", "nature", tmp_path / "x.docx")
    assert r.returncode == 2 and r.stderr.startswith("Error: manuscript dir not found")
    (manuscript / "results.md").write_bytes("结果很好。".encode("gbk"))
    r = run_export(manuscript, "nature", tmp_path / "x.docx")
    assert r.returncode == 2 and "Traceback" not in r.stderr and "not UTF-8" in r.stderr
    empty = tmp_path / "empty"
    empty.mkdir()
    r = run_export(empty, "nature", tmp_path / "x.docx")
    assert r.returncode == 2 and "no manuscript sections found" in r.stderr and "Traceback" not in r.stderr
    r = run_export(manuscript, "nature", tmp_path / "x.docx", "--no-such-flag")
    assert r.returncode == 2 and len(r.stderr.strip().splitlines()) == 1
