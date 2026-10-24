#!/usr/bin/env python3
"""
reproduce_check.py -- run an analysis several times in fresh processes, emptying its output
paths before every run, and check that it writes the same output files every time
(statistical-analysis Step 5, self-check). Files outside the output paths (a cache) are not
reset; see 2.

Usage
  python3 reproduce_check.py --cmd "python3 analysis_script.py" --outputs results/ tables/results.csv
  python3 reproduce_check.py --outputs results/ --runs 3 --report reproduce-check.md -- Rscript analysis.R

What it does
  1. Files that already exist in the output paths are MOVED (never deleted) to
     <keep-dir>/pre-existing/ so that run 1 starts with empty output paths. Directories that
     existed before the check (the output directories and their subdirectories) stay in
     place, emptied, because many scripts expect them to exist. After the last run, every
     pre-existing file that no run wrote again (a hand-made table, an output the analysis no
     longer writes) is copied back to its place. One that an earlier run wrote but the last run
     did not stays in pre-existing/ (putting the old version back would mix it with the new
     outputs); the report names it, and --keep-last keeps that check folder while it holds
     the only copy (unless an earlier check's run wrote that version). The same holds for a
     pre-existing file that the runs wrote again with other content (a hand-corrected table).
  2. The command runs in a fresh subprocess (--cwd, --timeout; empty stdin; stdout and stderr
     go to <keep-dir>/logs/); whatever it leaves running in its process group is stopped when it
     exits. After every run except the last, its outputs are moved to <keep-dir>/run<k>/ (and
     the directories it created for them removed), so every run starts from the same state.
     The last run's outputs stay where the script wrote them; a copy goes to <keep-dir>/run<N>/.
     Only --outputs is reset between runs: files the runs write elsewhere in --cwd (a cache)
     are listed in the report, because a later run may just read them back; so are files
     named like a cache (cache/, *_cache/, _targets/, _freeze/, *.pkl ...) that were there before
     the check, because every run reads the same copy. A cache with another name that predates
     the check looks like any input file and is not recognised: put it in --outputs. With more
     than 100000 files under --cwd only the names are checked, and the report says so.
  3. The runs are compared: the set of files and the sha256 of every file; CSV/TSV cell by
     cell; JSON / JSON Lines value by value; .txt/.md/.log/.Rout etc. (and any other file that
     is plain text) as text (the first 20 changed lines are shown); .xlsx cell values and
     formulas; .docx/.pptx by their content (not the saved time); .svg as text, ignoring
     matplotlib's embedded date and random element ids; .gz by its decompressed content;
     symlinks by where they point (never followed). Other files (PDF, PNG, RDS ...) by sha256
     only. Whole numbers and zero-padded codes ('007') must match exactly, and so must numbers
     that are the time of the runs; other numbers within --rtol / --atol. --exclude leaves files
     out of the comparison (a shell glob: * and ? stay within one directory, ** spans several).
  4. What the command prints is compared with --compare-stdout / --compare-stderr; without them
     the report still warns when the printed output differs between runs.
  5. If a run fails, or the check is interrupted (Ctrl-C, SIGTERM, SIGHUP), the analysis is
     stopped, its partial outputs go to <keep-dir>/run<k>-failed/ (-interrupted/) and the
     previous outputs are copied back into place.

--outputs, --report, --json and --keep-dir paths are relative to --cwd. <keep-dir> defaults to
<cwd>/.reproduce-check/<time>/; that folder gets its own .gitignore (the copies may hold
patient-level results), and --keep-last N deletes all but the N most recent checks -- only
folders this tool made (status file), and never one whose pre-existing/ holds a file whose
content is not also in its original place or in a check that is kept (so a hand-made or
hand-edited version is never deleted). The exception: a file that was an earlier check's result
(the same content as that check's last-run copy when this check started, which the status file
records) does not keep a folder, so an output that changes between checks does not leave one
folder behind for every second check.
Exit code: 0 identical (within tolerance) · 1 outputs differ · 2 a run failed, the check was
interrupted, or there was nothing to compare (also: invalid arguments).
"""

import argparse
import csv
import datetime as _dt
import difflib
import fnmatch
import gzip
import hashlib
import html
import io
import json
import math
import os
import re
import shlex
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from pathlib import Path

TABLE_SEPS = {".csv": ",", ".tsv": "\t", ".tab": "\t"}
TEXT_EXTS = {".txt", ".md", ".log", ".tex", ".html", ".htm", ".xml", ".yaml", ".yml",
             ".rmd", ".qmd", ".out", ".rout", ".smcl", ".lst", ".rtf"}
JSONL_EXTS = {".jsonl", ".ndjson"}
OFFICE_EXTS = {".docx", ".pptx"}                     # zip + XML: compared by content
EMBEDDED_TIME_EXTS = {".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
SCRIPT_EXTS = {".py", ".r", ".rmd", ".qmd", ".ipynb", ".jl", ".sh", ".bash", ".bat", ".cmd", ".ps1",
               ".do", ".ado", ".sas", ".sps", ".m", ".js"}
# --report / --json must never overwrite a script or a data file
NOT_A_REPORT_EXTS = SCRIPT_EXTS | {".csv", ".tsv", ".tab", ".xlsx", ".xlsm", ".xls", ".sav", ".dta", ".rds",
                                   ".rdata", ".rda", ".parquet", ".feather", ".sas7bdat"}
_NUM = r"[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?"
NUM_RE = re.compile(_NUM)
PLAIN_NUM_RE = re.compile(rf"^\s*{_NUM}\s*$")
INT_RE = re.compile(r"^\s*[-+]?\d+\s*$")
LEADING_ZERO_RE = re.compile(r"^\s*[-+]?0\d")         # '007', '0123': a code, not a quantity
MAX_LINES = 20
MAX_GZ = 256 << 20                                   # decompressed size limit for .gz / .docx outputs
TOOL = "reproduce_check.py"
MARKER = ".reproduce-check-run.json"                 # status file inside every check folder
FROM_CHECKS = "pre_existing_from_checks"             # status-file key: {rel: digest} of the pre-existing/
                                                     # files an earlier check's run wrote
CHECK_DIR_RE = re.compile(r"^\d{8}-\d{6}-")
CHECK_PARTS_RE = re.compile(r"^(?:pre-existing|logs|run\d+(?:-failed|-interrupted)?|\.reproduce-check-run\.json)$")
RUN_DIR_RE = re.compile(r"^run\d+(?:-failed|-interrupted)?$")
GITIGNORE = ("# reproduce_check.py: copies of analysis outputs from each check (they may contain\n"
             "# patient-level results). Never commit them; delete a check folder once it is reviewed.\n*\n")
TREE_SKIP = {".git", ".hg", ".svn", "__pycache__", "node_modules", ".ipynb_checkpoints", ".venv", "venv", "renv",
             "packrat", "site-packages", "dist-packages", ".Rproj.user", ".pytest_cache", ".mypy_cache", ".ruff_cache",
             ".tox"}
TREE_MAX = 100000                                    # files under --cwd compared for changes outside --outputs
TREE_WALK_MAX = 1000000                              # files looked at for cache-like names
# directories (or file stems) named like a cache: cache/, .cache/, report_cache/ (knitr), cache_dir/,
# cachedir/ and joblib/ (joblib.Memory), _targets/ (targets), .drake/, _freeze/ (Quarto), http_cache.sqlite
CACHE_NAME_RE = re.compile(r"(?i)^(?:[._]?caches?|.+[_.-]cache|cache[_.-].+|cachedir|joblib|_targets|\.drake|_freeze)$")
CACHE_EXTS = {".pkl", ".pickle", ".joblib", ".cache"}
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")

HINTS = (
    "未固定随机种子：Python 的 random.seed / numpy.random.default_rng(seed) / sklearn 的 random_state / "
    "torch.manual_seed，R 的 set.seed()——bootstrap、多重插补、交叉验证、机器学习都用到随机数",
    "输出依赖当前时间：结果里写了运行日期时间、文件名带时间戳；PDF 会内嵌保存时间（XLSX/DOCX/PPTX 已按内容比较，"
    "不受保存时间影响；matplotlib：PDF 用 savefig(..., metadata={'CreationDate': None})，SVG 用 metadata={'Date': None} "
    "并设 plt.rcParams['svg.hashsalt'] = 'fixed'；或设环境变量 SOURCE_DATE_EPOCH）；.gz 文件头里有压缩时间"
    "（pandas 用 compression={'method': 'gzip', 'mtime': 0}）",
    "多线程或 GPU 的非确定性：并行求和的顺序不同会带来末位差异；可设 OMP_NUM_THREADS=1、"
    "OPENBLAS_NUM_THREADS=1 重跑验证，确属浮点末位差异时再适当放宽 --rtol",
    "读取了会变的文件：运行中下载数据、读取上游仍在更新的文件、把上一次的输出当成输入",
    "顺序不固定：遍历集合（set）、os.listdir / glob 的结果没有排序",
    "确认某个文件只是内嵌的时间不同、内容相同后，可以用 --exclude 把它排除在比较之外，并在 analysis-log.md 里写明",
)


# ─── files ──────────────────────────────────────────────────────────────────

class Output:
    """One --outputs entry: a file or a directory, located relative to --cwd."""

    def __init__(self, raw, cwd):
        p = Path(raw).expanduser()
        self.raw = raw
        self.path = (p if p.is_absolute() else cwd / p).resolve()
        try:
            self.rel = self.path.relative_to(cwd).as_posix()
        except ValueError:
            self.rel = "_outside_cwd/" + "/".join(self.path.parts[1:])


def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _digest(path):
    """sha256 of a file; a symlink is identified by where it points (it is never followed)."""
    if os.path.islink(path):
        return "link:" + hashlib.sha256(os.readlink(path).encode("utf-8", "surrogateescape")).hexdigest()
    return _sha256(path)


def _files_at(root):
    """{sub_path: Path} for a file ('' key) or every file under a directory. Symlinks count as
    entries of their own and are never followed (a link to a data folder is not copied in)."""
    if root.is_symlink() or root.is_file():
        return {"": root}
    if not root.is_dir():
        return {}
    found = {}
    for d, dirs, files in os.walk(root):
        for name in files + [x for x in dirs if os.path.islink(os.path.join(d, x))]:
            p = Path(d) / name
            found[p.relative_to(root).as_posix()] = p
    return dict(sorted(found.items()))


def _collect(outputs, base=None):
    """All output files, keyed by their path relative to --cwd (base=None: the live location)."""
    found = {}
    for o in outputs:
        loc = o.path if base is None else base / o.rel
        for sub, p in _files_at(loc).items():
            found[o.rel + ("/" + sub if sub else "")] = p
    return found


def _is_real_dir(p):
    return p.is_dir() and not p.is_symlink()


def _count_files(p):
    return len(_files_at(p)) if _is_real_dir(p) else 1


def _dirs_under(outputs):
    """Every directory (not symlink) that is an output directory or lies inside one."""
    found = set()
    for o in outputs:
        if _is_real_dir(o.path):
            found.add(o.path)
            for root, dirs, _ in os.walk(o.path):
                found.update(Path(root) / d for d in dirs if not (Path(root) / d).is_symlink())
    return found


def _missing_parents(outputs, cwd):
    """Directories between --cwd and each output that do not exist yet (a run will create them)."""
    found = set()
    for o in outputs:
        for d in o.path.parents:
            if d == cwd or cwd not in d.parents:
                break
            if not os.path.lexists(d):
                found.add(d)
    return sorted(found, key=lambda d: len(d.parts), reverse=True)


def _remove_new_dirs(dirs):
    """Remove the (now empty) directories a run created around its outputs, deepest first; a
    directory that still holds something is left alone."""
    for d in dirs:
        try:
            os.rmdir(d)
        except OSError:
            pass


def _move_out(outputs, dest, keep_dirs=None):
    """Move current outputs into dest/<rel>. Returns the number of files moved.

    Directories in keep_dirs (None = every directory) stay in place, emptied; everything else --
    files, symlinks and directories a run created -- is moved, so the next run starts from the
    same state as run 1."""
    n = 0
    for o in outputs:
        if _is_real_dir(o.path) and (keep_dirs is None or o.path in keep_dirs):
            n += _move_children(o.path, dest / o.rel, keep_dirs)
        elif o.path.exists() or o.path.is_symlink():
            target = dest / o.rel
            target.parent.mkdir(parents=True, exist_ok=True)
            n += _count_files(o.path)
            shutil.move(str(o.path), str(target))
    return n


def _move_children(directory, target, keep_dirs):
    n = 0
    target.mkdir(parents=True, exist_ok=True)
    for child in sorted(directory.iterdir()):
        if _is_real_dir(child) and (keep_dirs is None or child in keep_dirs):
            n += _move_children(child, target / child.name, keep_dirs)
        else:
            n += _count_files(child)
            shutil.move(str(child), str(target / child.name))
    return n


def _copy_entry(src, dst, stop):
    """Copy one file, or one symlink as a symlink. Nothing is written through a symlink: not at
    dst, and not in a directory between `stop` and dst."""
    for d in dst.parents:
        if d == stop or stop not in d.parents:
            break
        if os.path.islink(d):
            raise OSError(f"{d} 是符号链接，不往里面写")
    dst.parent.mkdir(parents=True, exist_ok=True)
    if os.path.islink(dst) or (dst.exists() and not dst.is_dir()):
        dst.unlink()
    if os.path.islink(src):
        os.symlink(os.readlink(src), dst)
    else:
        shutil.copy2(src, dst)


def _copy_tree(src, dst, stop):
    """Copy a file or a directory tree; symlinks inside are copied as links, never followed."""
    if os.path.islink(src) or src.is_file():
        _copy_entry(src, dst, stop)
        return
    for sub, p in _files_at(src).items():
        _copy_entry(p, dst / sub, stop)
    for d, dirs, _ in os.walk(src):                       # empty directories too
        for name in dirs:
            if not os.path.islink(os.path.join(d, name)):
                (dst / Path(d).relative_to(src) / name).mkdir(parents=True, exist_ok=True)


def _copy_out(outputs, dest):
    for o in outputs:
        if o.path.exists() or o.path.is_symlink():
            _copy_tree(o.path, dest / o.rel, dest)


def _copy_back(outputs, src, cwd):
    """Copy a snapshot (src/<rel>) back to the live locations."""
    for o in outputs:
        snap = src / o.rel
        if _is_real_dir(snap):
            o.path.mkdir(parents=True, exist_ok=True)
        if snap.exists() or snap.is_symlink():
            _copy_tree(snap, o.path, cwd)


def _restore_missing(pre, cwd, snapshots):
    """After the last run: copy back every file from pre-existing/ that no run wrote again (a
    hand-made file, an output the analysis no longer writes). -> (restored, [{"file", "why"}] not
    put back): why is "dir" (the runs put a directory there), "earlier_run" (an earlier run wrote
    the file but the last one did not: the old version would sit among the new outputs) or the
    error that stopped the copy."""
    restored, kept = [], []
    if not _is_real_dir(pre):
        return restored, kept
    for rel, src in _files_at(pre).items():
        live = cwd / rel
        if os.path.lexists(live):
            if _is_real_dir(live):
                kept.append({"file": rel, "why": "dir"})
            continue
        if any(os.path.lexists(s / rel) and not _is_real_dir(s / rel) for s in snapshots):
            kept.append({"file": rel, "why": "earlier_run"})
            continue
        try:
            _copy_entry(src, live, cwd)
            restored.append(rel)
        except OSError as e:
            kept.append({"file": rel, "why": f"{type(e).__name__}: {e}"})
    return restored, kept


def _same_content(a, b, memo):
    """Do two entries hold the same content (a symlink: the same target)? Unreadable = different.
    memo caches sha256 by path for one pruning pass."""
    try:
        if os.path.islink(a) or os.path.islink(b):
            return os.path.islink(a) and os.path.islink(b) and os.readlink(a) == os.readlink(b)
        if not (os.path.isfile(a) and os.path.isfile(b)) or os.path.getsize(a) != os.path.getsize(b):
            return False
        for p in (a, b):
            if str(p) not in memo:
                memo[str(p)] = _sha256(p)
        return memo[str(a)] == memo[str(b)]
    except OSError:
        return False


def _digest_memo(path, memo):
    """_digest, with the sha256 of a regular file cached in memo (the same cache _same_content uses)."""
    if os.path.islink(path):
        return _digest(path)
    if str(path) not in memo:
        memo[str(path)] = _sha256(path)
    return memo[str(path)]


def _run_dirs(folders, pre_existing=False):
    """The run folders (run<k>, run<k>-failed, run<k>-interrupted) of these check folders (only
    those this tool made); with pre_existing=True their pre-existing/ folders too."""
    found = []
    for folder in folders:
        if not _made_here(_read_marker(folder)):
            continue
        try:
            found += [c for c in folder.iterdir() if _is_real_dir(c) and
                      (RUN_DIR_RE.match(c.name) or (pre_existing and c.name == "pre-existing"))]
        except OSError:
            continue
    return found


def _last_run_dir(folder):
    """The copy of what a check left in place: its highest-numbered run<k>/ (a run that completed)."""
    best, k = None, -1
    try:
        children = list(folder.iterdir())
    except OSError:
        return None
    for c in children:
        m = re.match(r"^run(\d+)$", c.name)
        if m and int(m.group(1)) > k and _is_real_dir(c):
            best, k = c, int(m.group(1))
    return best


def _from_earlier_checks(pre, current, keep_root, memo):
    """Files just moved into pre-existing/ that an earlier check's run wrote: the same content as that
    check's last-run copy (what it left in place), or as a file its status file records as one (a
    failed check copies its pre-existing/ back into place). -> {rel: digest}. Kept in the status file:
    --keep-last does not keep a check folder for the previous check's results; it protects only the
    other files (made or edited by hand, or written by a run outside the checks)."""
    found = {}
    files = _files_at(pre) if _is_real_dir(pre) else {}
    if not files:
        return found
    sources = []
    for folder, _ in _check_folders(keep_root):
        marker = _read_marker(folder)
        if folder == current or not _made_here(marker):
            continue
        rec = marker.get(FROM_CHECKS)
        sources.append((_last_run_dir(folder), rec if isinstance(rec, dict) else {}))
    for rel, src in files.items():
        for last, rec in reversed(sources):                # the most recent check first
            try:
                if (last is not None and os.path.lexists(last / rel) and _same_content(src, last / rel, memo)) \
                        or (rel in rec and rec[rel] == _digest_memo(src, memo)):
                    found[rel] = _digest_memo(src, memo)
                    break
            except OSError:
                continue
    return found


def _only_copies(pre, base, copy_dirs, memo, record=None):
    """Files in pre-existing/ whose content exists nowhere else: -> (gone, changed). gone: their
    place is empty (or a directory); changed: a file with other content is there now (e.g. a
    hand-edited table the runs wrote again). A file is not counted when the same content is in its
    place, or in one of copy_dirs (a run folder -- written by a check's run: an analysis output, not
    a hand-made file -- or the pre-existing/ of a check that is kept), or when record (from the
    check's status file) says an earlier check's run wrote it and it still has that content."""
    gone, changed = [], []
    if not _is_real_dir(pre):
        return gone, changed
    for rel, src in _files_at(pre).items():
        if not rel:
            continue
        if record and rel in record:
            try:
                if record[rel] == _digest_memo(src, memo):
                    continue
            except OSError:
                pass
        live = base / rel
        in_place = os.path.lexists(live) and not _is_real_dir(live)
        if in_place and _same_content(src, live, memo):
            continue
        if any(os.path.lexists(d / rel) and _same_content(src, d / rel, memo) for d in copy_dirs):
            continue
        (changed if in_place else gone).append(rel)
    return gone, changed


def _glob_parts(pat_parts, parts):
    """Shell-style match of path segments: * ? [..] stay inside one segment, ** spans any number."""
    if not pat_parts:
        return not parts
    if pat_parts[0] == "**":
        return any(_glob_parts(pat_parts[1:], parts[i:]) for i in range(len(parts) + 1))
    return bool(parts) and fnmatch.fnmatchcase(parts[0], pat_parts[0]) and _glob_parts(pat_parts[1:], parts[1:])


def _excluded(rel, patterns):
    """True if an output file (path relative to --cwd) matches an --exclude pattern: a glob on the
    path relative to --cwd (* and ? do not cross '/', ** spans directories), a name or glob without
    '/' on the file name, or a directory (trailing '/' optional) that the file lies in."""
    parts = rel.split("/")
    for pat in patterns or ():
        p = pat.replace("\\", "/")
        while p.startswith("./"):
            p = p[2:]
        pp = [x for x in p.split("/") if x]
        if not pp:
            continue
        if not p.endswith("/") and ((len(pp) == 1 and fnmatch.fnmatchcase(parts[-1], pp[0]))
                                    or _glob_parts(pp, parts)):
            return True
        if any(_glob_parts(pp, parts[:i]) for i in range(1, len(parts))):
            return True
    return False


def _norm_pattern(pat, cwd):
    """An absolute --exclude pattern inside --cwd is made relative to it."""
    if os.path.isabs(pat):
        for p in (os.path.normpath(pat), os.path.realpath(pat)):     # realpath: /tmp -> /private/tmp
            try:
                return Path(p).relative_to(cwd).as_posix() + ("/" if pat.endswith(("/", "\\")) else "")
            except ValueError:
                continue
    return pat


def _cache_like(rel, dir_hit=None):
    """True for a path named like a cache: a directory (or the file's stem) such as cache/,
    report_cache/, _targets/, or a pickled object (.pkl, .joblib). dir_hit: whether a directory
    on the path is already known to be named like one."""
    d, _, name = rel.rpartition("/")
    stem, ext = os.path.splitext(name)
    ext = ext.lower()
    if ext in SCRIPT_EXTS:                                 # code (a module called cache.py) is not a cache
        return False
    if dir_hit is None:
        dir_hit = any(CACHE_NAME_RE.match(x) for x in d.split("/") if x)
    return dir_hit or ext in CACHE_EXTS or bool(CACHE_NAME_RE.match(stem))


def _tree_state(cwd, skip):
    """Files under --cwd outside the outputs and the check folders, to spot state the runs leave
    for each other (a cache). -> {"files": {rel: (size, mtime, atime)} or None when there are more
    than TREE_MAX files, "cache_like": [rel], "complete": False when the walk stopped early}."""
    state, cache, n = {}, [], 0
    out = {"files": state, "cache_like": cache, "complete": True}
    base = str(cwd)
    skip = {str(p) for p in skip}
    try:
        for root, dirs, files in os.walk(base):
            # a Python / conda environment, whatever its name -- but not --cwd itself (a project with
            # `python -m venv .` in it still has its cache to find)
            if root != base and ("pyvenv.cfg" in files or "conda-meta" in dirs):
                dirs[:] = []
                continue
            dirs[:] = sorted(d for d in dirs if d not in TREE_SKIP and os.path.join(root, d) not in skip)
            rel_root = os.path.relpath(root, base).replace(os.sep, "/")
            prefix = "" if rel_root == "." else rel_root + "/"
            dir_hit = any(CACHE_NAME_RE.match(x) for x in prefix.split("/") if x)
            for f in files:
                p = os.path.join(root, f)
                if p in skip:
                    continue
                n += 1
                if n > TREE_WALK_MAX:
                    out["complete"] = False
                    return out
                rel = prefix + f
                if _cache_like(rel, dir_hit):
                    cache.append(rel)
                if out["files"] is None:
                    continue
                try:
                    st = os.lstat(p)
                except OSError:
                    continue
                state[rel] = (st.st_size, st.st_mtime_ns, st.st_atime_ns)
                if len(state) > TREE_MAX:
                    out["files"] = None                    # too many to compare: cache-like names only
    except OSError:
        out["complete"] = False
    return out


# ─── interruption ───────────────────────────────────────────────────────────

class Interrupted(Exception):
    """Ctrl-C, SIGTERM or SIGHUP arrived while the check was running."""

    def __init__(self, signum):
        super().__init__(signum)
        self.signum = signum


_SIG = {"defer": 0, "pending": None, "keep": None}   # keep: the check folder once files may move ("done": reported)


def _on_signal(signum, frame):
    if _SIG["defer"]:
        _SIG["pending"] = signum          # acted on as soon as the current file move is complete
        return
    raise Interrupted(signum)


class _Deferred:
    """Signals that arrive while files are being moved are acted on right after the move."""

    def __enter__(self):
        _SIG["defer"] += 1

    def __exit__(self, exc_type, exc, tb):
        _SIG["defer"] -= 1
        if not _SIG["defer"] and _SIG["pending"] is not None and exc_type is None:
            signum, _SIG["pending"] = _SIG["pending"], None
            raise Interrupted(signum)
        return False


def _install_handlers():
    old = {}
    if threading.current_thread() is not threading.main_thread():
        return old
    for name in ("SIGINT", "SIGTERM", "SIGHUP"):
        sig = getattr(signal, name, None)
        if sig is None:
            continue
        try:
            old[sig] = signal.signal(sig, _on_signal)
        except (ValueError, OSError):
            pass
    return old


def _restore_handlers(old):
    for sig, handler in old.items():
        try:
            signal.signal(sig, handler)
        except (ValueError, OSError, TypeError):
            pass


def _signal_name(signum):
    try:
        return signal.Signals(signum).name
    except (ValueError, TypeError):
        return str(signum)


# ─── running ────────────────────────────────────────────────────────────────

def _kill(proc):
    """Stop the command and everything it started (not only the shell)."""
    try:
        if os.name == "posix":
            os.killpg(proc.pid, signal.SIGKILL)       # the whole process group (shell + children)
        else:                                          # Windows: kill the process tree
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)], stdout=subprocess.DEVNULL,
                           stderr=subprocess.DEVNULL, timeout=30)
    except (OSError, subprocess.SubprocessError):
        pass
    try:
        proc.kill()
    except OSError:
        pass
    try:
        proc.wait(timeout=10)
    except subprocess.TimeoutExpired:
        pass


def _group_alive(pgid):
    """Is a process of this process group (not a zombie) still running?"""
    proc_dir = Path("/proc")
    try:
        entries = [e for e in os.listdir(proc_dir) if e.isdigit()]
    except OSError:
        entries = None
    if entries is not None:
        for e in entries:
            try:
                stat = (proc_dir / e / "stat").read_text(encoding="utf-8", errors="replace")
            except OSError:
                continue
            fields = stat[stat.rfind(")") + 2:].split()       # state ppid pgrp ...
            if len(fields) > 2 and fields[2] == str(pgid) and fields[0] not in "ZX":
                return True
        return False
    try:
        os.killpg(pgid, 0)
        return True
    except OSError:
        return False


def _end_group(proc):
    """After the command exited: stop what it left running in its process group (a background job
    would otherwise write into the next run's outputs). True if something was still running."""
    if os.name != "posix":
        return False
    alive = _group_alive(proc.pid)
    try:
        os.killpg(proc.pid, signal.SIGKILL)
    except OSError:
        return False
    return alive


def _read_log(path):
    try:
        return Path(path).read_bytes().decode("utf-8", "replace")
    except OSError:
        return ""


def run_once(cmd, shell, cwd, timeout, log_dir, k):
    """Run the command once. stdout / stderr go straight to files, so a process the command
    leaves behind cannot keep the checker waiting on a pipe."""
    log_dir.mkdir(parents=True, exist_ok=True)
    out_p, err_p = log_dir / f"run{k}.stdout.txt", log_dir / f"run{k}.stderr.txt"
    res = {"run": k, "returncode": None, "timed_out": False, "launch_error": None, "background": False}
    t0 = time.time()
    with open(out_p, "wb") as fo, open(err_p, "wb") as fe:
        kw = {"cwd": str(cwd), "stdin": subprocess.DEVNULL, "stdout": fo, "stderr": fe, "shell": shell}
        if os.name == "posix":
            kw["start_new_session"] = True
        else:
            kw["creationflags"] = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
        try:
            proc = subprocess.Popen(cmd, **kw)
        except OSError as e:
            res["launch_error"] = str(e)
            fe.write(str(e).encode("utf-8", "replace"))
        else:
            try:
                proc.wait(timeout=timeout)
                res["returncode"] = proc.returncode
                res["background"] = _end_group(proc)
            except subprocess.TimeoutExpired:
                _kill(proc)
                res["timed_out"] = True
            except BaseException:
                _kill(proc)                            # Ctrl-C / SIGTERM: never leave the analysis running
                raise
    res["seconds"] = round(time.time() - t0, 1)
    res["stdout"], res["stderr"] = _read_log(out_p), _read_log(err_p)
    res["ok"] = res["returncode"] == 0
    return res


def _signal_of(returncode, shell):
    """Name of the signal that ended the command, if the exit code says so."""
    if returncode is None:
        return None
    if returncode < 0:
        return _signal_name(-returncode)
    if shell and os.name == "posix" and 128 < returncode < 128 + 65:  # the shell reports 128 + signal
        name = _signal_name(returncode - 128)
        return name if name.startswith("SIG") else None
    return None


# ─── comparing ──────────────────────────────────────────────────────────────

_WINDOW = {"start": None, "end": None}               # when the runs happened (epoch seconds)


def _is_run_time(x):
    """True if x is a moment during the runs, in seconds / ms / µs / ns since 1970: a timestamp."""
    s, e = _WINDOW["start"], _WINDOW["end"]
    if s is None or e is None or not isinstance(x, (int, float)) or isinstance(x, bool):
        return False
    return any(s - 1 <= x / scale <= e + 1 for scale in (1, 1e3, 1e6, 1e9))


def _raise_csv_limit():
    limit = sys.maxsize
    while True:
        try:
            csv.field_size_limit(limit)
            return
        except OverflowError:
            limit //= 10


def _bytes(src):
    return bytes(src) if isinstance(src, (bytes, bytearray)) else Path(src).read_bytes()


def _read_text(src):
    data = _bytes(src)
    for enc in ("utf-8-sig", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _looks_like_text(src):
    """True for a file (or bytes) that is plain UTF-8 text: the first 64 KB decode, with no NUL bytes."""
    try:
        if isinstance(src, (bytes, bytearray)):
            head = bytes(src[:65536])
        else:
            with open(src, "rb") as fh:
                head = fh.read(65536)
    except OSError:
        return False
    if b"\0" in head:
        return False
    try:
        head.decode("utf-8")
        return True
    except UnicodeDecodeError as e:
        return len(head) == 65536 and e.start >= len(head) - 3     # a character cut at the 64 KB mark


def _as_number(v):
    """int for whole numbers ('12', 12), float for other numbers, None for anything else."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return v
    if isinstance(v, str):
        if INT_RE.match(v):
            return int(v)
        if PLAIN_NUM_RE.match(v):
            return float(v)
    return None


def _num_state(a, b, rtol, atol):
    """-> 'same' | 'tol' | 'diff' for two numbers. Whole numbers (counts, IDs, epoch seconds) and
    timestamps of the runs must match exactly; other numbers may differ within the tolerance."""
    if isinstance(a, int) and isinstance(b, int):
        return "same" if a == b else "diff"
    try:
        fa, fb = float(a), float(b)
    except OverflowError:
        return "same" if a == b else "diff"
    if fa == fb or (math.isnan(fa) and math.isnan(fb)):
        return "same"
    if math.isinf(fa) or math.isinf(fb) or math.isnan(fa) or math.isnan(fb):
        return "diff"
    if _is_run_time(fa) and _is_run_time(fb):
        return "diff"                                 # the time of each run, whatever the tolerance
    return "tol" if math.isclose(fa, fb, rel_tol=rtol, abs_tol=atol) else "diff"


def _str_num_state(u, v, rtol, atol):
    """Two numbers written as text. Whole numbers and zero-padded codes ('007' is not '07') must be
    written the same; other numbers are compared with _num_state."""
    if u.strip() == v.strip():
        return "same"
    if (INT_RE.match(u) and INT_RE.match(v)) or LEADING_ZERO_RE.match(u) or LEADING_ZERO_RE.match(v):
        return "diff"
    return _num_state(_as_number(u), _as_number(v), rtol, atol)


def _ts_note(a, b):
    fa, fb = _as_number(a), _as_number(b)
    return "（像是运行时的时间戳）" if _is_run_time(fa) and _is_run_time(fb) else ""


class _Formula:
    """An xlsx formula cell: its formula text plus the value Excel cached for it (if any)."""
    __slots__ = ("text", "cached")

    def __init__(self, text, cached):
        self.text, self.cached = text, cached

    def __repr__(self):
        return f"{self.text}（缓存值 {self.cached!r}）"


def _cells_equal(va, vb, rtol, atol):
    """-> 'same' | 'tol' | 'diff' for two cell values (str / number / formula / other)."""
    if isinstance(va, _Formula) or isinstance(vb, _Formula):
        if not (isinstance(va, _Formula) and isinstance(vb, _Formula)) or va.text != vb.text:
            return "diff"
        return _cells_equal(va.cached, vb.cached, rtol, atol)
    if isinstance(va, bool) or isinstance(vb, bool):
        return "same" if type(va) is type(vb) and va == vb else "diff"
    if va == vb:
        return "same"
    if isinstance(va, str) and isinstance(vb, str):
        if not (PLAIN_NUM_RE.match(va) and PLAIN_NUM_RE.match(vb)):
            return "diff"
        state = _str_num_state(va, vb, rtol, atol)
        return "tol" if state == "same" else state    # the same number written differently ("1" / "1.0")
    if any(isinstance(v, str) and LEADING_ZERO_RE.match(v) for v in (va, vb)):
        return "diff"                                 # the code '007' is not the number 7
    fa, fb = _as_number(va), _as_number(vb)
    if fa is not None and fb is not None:
        state = _num_state(fa, fb, rtol, atol)
        return "tol" if state == "same" else state
    return "diff"


def _compare_grid(ra, rb, rtol, atol, where=""):
    details, n_diff, n_tol = [], 0, 0
    if len(ra) != len(rb):
        details.append(f"{where}行数不同：{len(ra)} → {len(rb)}")
        n_diff += abs(len(ra) - len(rb))
    header = [str(h) for h in ra[0]] if ra else []
    for i, (row_a, row_b) in enumerate(zip(ra, rb)):
        row_a, row_b = list(row_a), list(row_b)
        if len(row_a) != len(row_b):
            details.append(f"{where}第 {i + 1} 行列数不同：{len(row_a)} → {len(row_b)}")
            n_diff += 1
        for j, (va, vb) in enumerate(zip(row_a, row_b)):
            state = _cells_equal(va, vb, rtol, atol)
            if state == "tol":
                n_tol += 1
            elif state == "diff":
                n_diff += 1
                if len(details) < MAX_LINES:
                    col = header[j] if i > 0 and j < len(header) else f"第 {j + 1} 列"
                    details.append(f"{where}第 {i + 1} 行 [{col}]：{va!r} → {vb!r}{_ts_note(va, vb)}")
    return details, n_diff, n_tol


def compare_table(a, b, sep, rtol, atol):
    _raise_csv_limit()                     # a long cell (a JSON blob, a free-text note) is not an error
    ra = list(csv.reader(io.StringIO(_read_text(a)), delimiter=sep))
    rb = list(csv.reader(io.StringIO(_read_text(b)), delimiter=sep))
    details, n_diff, n_tol = _compare_grid(ra, rb, rtol, atol)
    return {"method": "逐单元格比较", "equal": n_diff == 0, "n_diff": n_diff, "n_within_tol": n_tol,
            "details": details}


def _xlsx_sheets(src, openpyxl):
    """{sheet: rows}; a formula cell becomes _Formula(formula text, cached value)."""
    data = _bytes(src)
    wf = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    wv = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = {}
        for name in wf.sheetnames:
            f_rows = [list(r) for r in wf[name].iter_rows(values_only=True)]
            v_rows = [list(r) for r in wv[name].iter_rows(values_only=True)]
            grid = []
            for i, row in enumerate(f_rows):
                vrow = v_rows[i] if i < len(v_rows) else []
                grid.append([_Formula(c, vrow[j] if j < len(vrow) else None)
                             if isinstance(c, str) and c.startswith("=") else c for j, c in enumerate(row)])
            sheets[name] = grid
        return list(wf.sheetnames), sheets
    finally:
        wf.close()
        wv.close()


def compare_xlsx(a, b, rtol, atol):
    try:
        import openpyxl
    except ImportError:
        res = compare_office(a, b, rtol, atol, "xlsx", texts=False)
        res["method"] = "xlsx 按 zip 内容比较（未安装 openpyxl，无法逐单元格比较；忽略内嵌保存时间）"
        return res
    names_a, sa = _xlsx_sheets(a, openpyxl)
    names_b, sb = _xlsx_sheets(b, openpyxl)
    details, n_diff, n_tol = [], 0, 0
    if names_a != names_b:
        details.append(f"工作表不同：{names_a} → {names_b}")
        n_diff += 1
    for name in names_a:
        if name not in sb:
            continue
        d, nd, nt = _compare_grid(sa[name], sb[name], rtol, atol, f"[{name}] ")
        details += d
        n_diff += nd
        n_tol += nt
    return {"method": "逐单元格比较（xlsx 公式与取值，忽略内嵌保存时间）", "equal": n_diff == 0, "n_diff": n_diff,
            "n_within_tol": n_tol, "details": details[:MAX_LINES]}


_OFFICE_TIME = re.compile(r"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)")
_OFFICE_RUN = re.compile(r"<(?:w|a):t(?:\s[^>]*)?>([^<]*)</(?:w|a):t>")
_OFFICE_PARA = re.compile(r"</(?:w|a):p>")


def _zip_parts(src):
    """{member: bytes} of a zip-based document (DOCX / PPTX / XLSX); the saved times in
    docProps/core.xml are blanked. The zip's own file dates are not part of the content."""
    with zipfile.ZipFile(io.BytesIO(_bytes(src))) as z:
        infos = [i for i in z.infolist() if not i.is_dir()]
        if sum(i.file_size for i in infos) > MAX_GZ:
            raise ValueError("解压后超过 256 MB，只按 sha256 比较")
        parts = {i.filename: z.read(i) for i in infos}
    core = parts.get("docProps/core.xml")
    if core is not None:
        parts["docProps/core.xml"] = _OFFICE_TIME.sub(r"\1\2", core.decode("utf-8", "replace")).encode("utf-8")
    return parts


def _office_text(xml):
    """The visible text of a Word / PowerPoint XML part, one paragraph per line."""
    return "\n".join("".join(html.unescape(t) for t in _OFFICE_RUN.findall(p))
                     for p in _OFFICE_PARA.split(xml)).strip()


def compare_office(a, b, rtol, atol, label, texts=True):
    pa, pb = _zip_parts(a), _zip_parts(b)
    details, n_diff = [], 0
    for name in sorted(set(pa) ^ set(pb)):
        details.append(f"{name}：只在其中一次运行中出现")
        n_diff += 1
    for name in sorted(set(pa) & set(pb)):
        if pa[name] == pb[name]:
            continue
        if not name.endswith((".xml", ".rels")):
            details.append(f"{name}：内容不同（嵌入的图片或其他文件）")
            n_diff += 1
            continue
        if compare_text(pa[name], pb[name], rtol, atol)["equal"]:
            continue                                   # only numbers' last digits, within the tolerance
        n_diff += 1
        xa, xb = _office_text(_read_text(pa[name])), _office_text(_read_text(pb[name]))
        if not texts:
            details.append(f"{name}：内容不同")
        elif xa != xb:
            details += [f"{name}：{x}" for x in compare_text(xa.encode("utf-8"), xb.encode("utf-8"), rtol, atol)["details"]]
        else:
            details.append(f"{name}：文字相同，格式或版式不同")
    method = f"{label} 按内容比较（忽略 zip 里的文件时间和文档属性里的保存时间）"
    if n_diff == 0:
        return {"method": method, "equal": True, "details": ["内容相同，只有保存时间（或在容差内的数字末位）不同"]}
    return {"method": method, "equal": False, "n_diff": n_diff, "details": details[:MAX_LINES]}


def _walk_json(x, y, path, out, rtol, atol):
    if len(out) >= 200:
        return
    if isinstance(x, bool) or isinstance(y, bool) or x is None or y is None:
        if not (type(x) is type(y) and x == y):          # true is not 1, null is not 0
            out.append(f"{path}: {x!r} → {y!r}")
    elif isinstance(x, (int, float)) and isinstance(y, (int, float)):
        if _num_state(x, y, rtol, atol) == "diff":
            out.append(f"{path}: {x!r} → {y!r}{_ts_note(x, y)}")
    elif isinstance(x, dict) and isinstance(y, dict):
        for k in sorted(set(x) | set(y), key=str):
            if k not in x or k not in y:
                out.append(f"{path}.{k}: 只在其中一次运行中出现")
            else:
                _walk_json(x[k], y[k], f"{path}.{k}", out, rtol, atol)
    elif isinstance(x, list) and isinstance(y, list):
        if len(x) != len(y):
            out.append(f"{path}: 列表长度 {len(x)} → {len(y)}")
        for i, (u, v) in enumerate(zip(x, y)):
            _walk_json(u, v, f"{path}[{i}]", out, rtol, atol)
    elif x != y:
        out.append(f"{path}: {x!r} → {y!r}")


def compare_text(a, b, rtol, atol, normalise=None):
    ta, tb = _read_text(a), _read_text(b)
    if normalise:
        ta, tb = normalise(ta), normalise(tb)
    if ta == tb:
        return {"method": "文本", "equal": True, "details": []}
    if ta.replace("\r\n", "\n") == tb.replace("\r\n", "\n"):
        return {"method": "文本", "equal": False, "n_diff": 1, "details": ["只有换行符不同（\\r\\n 与 \\n）"]}
    na, nb = NUM_RE.findall(ta), NUM_RE.findall(tb)
    if NUM_RE.split(ta) == NUM_RE.split(tb) and len(na) == len(nb) and \
            all(_str_num_state(u, v, rtol, atol) != "diff" for u, v in zip(na, nb)):
        return {"method": "文本（数字按容差比较）", "equal": True, "details": ["只有数字的末位不同，在容差内"]}
    changed = [ln for ln in difflib.unified_diff(ta.splitlines(), tb.splitlines(), lineterm="", n=0)
               if ln[:1] in "+-" and not ln.startswith(("+++", "---"))]
    if not changed:
        changed = ["差别在看不见的字符上（行尾空白或文件末尾的换行）"]
    return {"method": "文本 diff", "equal": False, "n_diff": len(changed), "details": changed[:MAX_LINES]}


def compare_json(a, b, rtol, atol):
    try:
        ja, jb = json.loads(_read_text(a)), json.loads(_read_text(b))
    except ValueError:
        return compare_text(a, b, rtol, atol)
    diffs = []
    _walk_json(ja, jb, "$", diffs, rtol, atol)
    return {"method": "JSON 逐值比较（整数须完全相同，其他数字按容差）", "equal": not diffs, "n_diff": len(diffs),
            "details": diffs[:MAX_LINES]}


def compare_jsonl(a, b, rtol, atol):
    try:
        la = [json.loads(x) for x in _read_text(a).splitlines() if x.strip()]
        lb = [json.loads(x) for x in _read_text(b).splitlines() if x.strip()]
    except ValueError:
        return compare_text(a, b, rtol, atol)
    diffs = [f"行数不同：{len(la)} → {len(lb)}"] if len(la) != len(lb) else []
    for i, (x, y) in enumerate(zip(la, lb)):
        _walk_json(x, y, f"第 {i + 1} 行 $", diffs, rtol, atol)
    return {"method": "JSON Lines 逐行逐值比较（整数须完全相同，其他数字按容差）", "equal": not diffs,
            "n_diff": len(diffs), "details": diffs[:MAX_LINES]}


_SVG_DATE = re.compile(r"<dc:date>[^<]*</dc:date>")
# matplotlib names clip paths / markers / hatches with a letter + 10 random hex digits unless
# rcParams['svg.hashsalt'] is set
_SVG_ID = re.compile(r"(id=\"|url\(#|href=\"#)([A-Za-z])[0-9a-f]{10}\b")


def _svg_normalise(text):
    return _SVG_ID.sub(r"\1\2-id", _SVG_DATE.sub("<dc:date/>", text))


def compare_svg(a, b, rtol, atol):
    res = compare_text(a, b, rtol, atol, normalise=_svg_normalise)
    res["method"] = "SVG 文本（忽略内嵌日期与 matplotlib 的随机图元编号）"
    if res["equal"]:
        res["details"] = ["只有内嵌日期或随机编号不同；作图时可固定：savefig(..., metadata={'Date': None})，"
                          "plt.rcParams['svg.hashsalt'] = 'fixed'"]
    return res


def _gunzip(src):
    with gzip.GzipFile(fileobj=io.BytesIO(_bytes(src))) as fh:
        data = fh.read(MAX_GZ + 1)
    if len(data) > MAX_GZ:
        raise ValueError("解压后超过 256 MB，只按 sha256 比较")
    return data


def compare_gzip(a, b, rtol, atol, name):
    ra, rb = _gunzip(a), _gunzip(b)
    if ra == rb:
        return {"method": "gzip 解压后逐字节比较", "equal": True,
                "details": ["解压后的内容完全相同，只有 gzip 文件头（压缩时间）不同；"
                            "pandas 可用 compression={'method': 'gzip', 'mtime': 0} 去掉"]}
    res = compare_file(ra, rb, rtol, atol, name=name[:-3])
    return dict(res, method="gzip 解压后" + res["method"])


def _binary_note(ext):
    note = ["二进制文件内容不同，无法逐项比较"]
    if ext in EMBEDDED_TIME_EXTS:
        note.append("这类文件常内嵌生成时间：确认内容相同后，可让脚本去掉时间戳，或用 --exclude 排除并写明")
    elif ext in IMAGE_EXTS:
        note.append("图片不同多半是作图用的数据不同，或抖动（jitter）/布局用了没固定的随机数")
    return {"method": "sha256", "equal": False, "n_diff": None, "details": note}


def compare_file(a, b, rtol, atol, name=None):
    """Compare two versions of one output (paths or bytes; name gives the file type)."""
    name = str(name if name is not None else a)
    ext = Path(name).suffix.lower()
    if ext in TABLE_SEPS:
        kind, fn = "表格", lambda: compare_table(a, b, TABLE_SEPS[ext], rtol, atol)
    elif ext == ".json":
        kind, fn = "JSON", lambda: compare_json(a, b, rtol, atol)
    elif ext in JSONL_EXTS:
        kind, fn = "JSON Lines", lambda: compare_jsonl(a, b, rtol, atol)
    elif ext in (".xlsx", ".xlsm"):
        kind, fn = "xlsx", lambda: compare_xlsx(a, b, rtol, atol)
    elif ext in OFFICE_EXTS:
        kind, fn = ext[1:].upper(), lambda: compare_office(a, b, rtol, atol, ext[1:].upper())
    elif ext == ".svg":
        kind, fn = "SVG 文本", lambda: compare_svg(a, b, rtol, atol)
    elif ext in TEXT_EXTS:
        kind, fn = "文本", lambda: compare_text(a, b, rtol, atol)
    elif ext == ".gz":
        kind, fn = "gzip", lambda: compare_gzip(a, b, rtol, atol, name)
    elif ext not in EMBEDDED_TIME_EXTS and ext not in IMAGE_EXTS and _looks_like_text(a) and _looks_like_text(b):
        kind, fn = "文本", lambda: compare_text(a, b, rtol, atol)  # .Rout-like logs with an unlisted extension
    else:
        return _binary_note(ext)
    try:
        return fn()
    except Exception as e:                  # a damaged or mislabelled file is a difference, not a crash
        return {"method": f"sha256（无法按{kind}读取）", "equal": False, "n_diff": None,
                "details": [f"按{kind}读取时出错：{type(e).__name__}: {str(e)[:200]}",
                            "两次运行的文件内容不同（sha256 不一致）"]}


def _link_text(p):
    return f"链接 → {os.readlink(p)}" if os.path.islink(p) else "普通文件"


def compare_runs(snapshots, outputs, rtol, atol, exclude=()):
    maps = [_collect(outputs, base=s) for s in snapshots]
    names = sorted(set().union(*[set(m) for m in maps])) if maps else []
    excluded = [n for n in names if _excluded(n, exclude)]
    unmatched = [p for p in exclude or () if not any(_excluded(n, [p]) for n in names)]
    names = [n for n in names if n not in excluded]
    res = {"n_files": len(names), "identical": [], "within_tolerance": [], "different": [], "excluded": excluded,
           "unmatched_exclude": unmatched}
    for rel in names:
        missing = [i + 1 for i, m in enumerate(maps) if rel not in m]
        if missing:
            res["different"].append({"file": rel, "method": "文件集合", "equal": False,
                                     "details": [f"第 {'、'.join(map(str, missing))} 次运行没有生成这个文件"]})
            continue
        hashes = [_digest(m[rel]) for m in maps]
        if len(set(hashes)) == 1:
            res["identical"].append(rel)
            continue
        verdict = None
        for k in range(1, len(maps)):
            if hashes[k] == hashes[0]:
                continue
            a, b = maps[0][rel], maps[k][rel]
            if os.path.islink(a) or os.path.islink(b):
                cmp = {"method": "符号链接（比较指向，不跟随）", "equal": False, "runs": [1, k + 1],
                       "details": [f"{_link_text(a)}；第 {k + 1} 次：{_link_text(b)}"]}
            else:
                cmp = dict(compare_file(a, b, rtol, atol), runs=[1, k + 1])
            if verdict is None or not cmp["equal"]:
                verdict = cmp
            if not cmp["equal"]:
                break
        (res["within_tolerance"] if verdict["equal"] else res["different"]).append({"file": rel, **verdict})
    empty = [o.raw for o in outputs if not any(_files_at(s / o.rel) for s in snapshots)]
    res["never_produced"] = empty
    return res


def compare_stream(logs, rtol, atol, compared):
    """What the runs printed (stdout or stderr). Always checked; it only decides the verdict with
    --compare-stdout / --compare-stderr."""
    texts = [_read_log(p) for p in logs]
    res = {"compared": compared, "equal": True, "empty": not any(t.strip() for t in texts),
           "method": "文本", "details": [], "runs": None}
    for k in range(1, len(logs)):
        c = compare_text(logs[0], logs[k], rtol, atol)
        if not c["equal"]:
            res.update(equal=False, method=c["method"], details=c["details"], runs=[1, k + 1])
            break
    return res


def _date_regex(days, window=None):
    """The check's own date(s) in the usual spellings and, when the runs took at most three hours,
    the clock times (HH:MM) during the runs."""
    pats = []
    for d in sorted(days):
        y, m, dd = d.year, d.month, d.day
        mon = rf"(?i:{MONTHS[m - 1]}[a-z]*\.?)"
        pats += [rf"(?<!\d){y}[-/.]0?{m}[-/.]0?{dd}(?!\d)", rf"{y}\s*年\s*0?{m}\s*月\s*0?{dd}\s*日",
                 rf"(?<!\d)0?{dd}[-/.]0?{m}[-/.]{y}(?!\d)", rf"(?<!\d)0?{m}[-/.]0?{dd}[-/.]{y}(?!\d)",
                 rf"(?<!\d){y}{m:02d}{dd:02d}(?!\d)",                         # 20261024, 20261024_0321
                 rf"\b{mon}\s+0?{dd}\b[^\n]{{0,20}}?(?<!\d){y}(?!\d)",         # Oct 24, 2026 · Sat Oct 24 03:21:00 2026
                 rf"(?<!\d)0?{dd}[\s-]?{mon}[\s,-]*{y}(?!\d)"]                 # 24 Oct 2026 · 24OCT2026 (SAS)
    if window and window[0] is not None and window[1] - window[0] <= 3 * 3600:
        stamps, t = set(), int(window[0] // 60 * 60)
        while t <= window[1]:
            stamps.add(time.strftime("%H:%M", time.localtime(t)))
            t += 60
        pats += [rf"(?<![\d:]){s}(?!\d)" for s in sorted(stamps)]
    return re.compile("|".join(pats))


def _scan_text(rel, p):
    """Text of an output file for the run-date check (None for binary files and symlinks)."""
    if os.path.islink(p):
        return None
    ext = Path(rel).suffix.lower()
    try:
        if ext == ".gz":
            with gzip.open(p, "rb") as fh:
                data = fh.read(2 << 20)
            return _read_text(data) if _looks_like_text(data) else None
        if ext in OFFICE_EXTS:
            with zipfile.ZipFile(p) as z:
                xml = [i for i in z.infolist() if re.match(r"^(word/document|ppt/slides/slide\d+)\.xml$", i.filename)
                       and i.file_size <= 8 << 20]
                return "\n".join(_office_text(z.read(i).decode("utf-8", "replace")) for i in xml[:200])
        if ext in EMBEDDED_TIME_EXTS or ext in IMAGE_EXTS or ext in (".xlsx", ".xlsm"):
            return None
        with open(p, "rb") as fh:
            head = fh.read(2 << 20)
        if ext in TEXT_EXTS or ext in TABLE_SEPS or ext in JSONL_EXTS or ext in (".json", ".svg") \
                or _looks_like_text(head):
            text = _read_text(head)
            return _svg_normalise(text) if ext == ".svg" else text
    except Exception:                        # unreadable: nothing to scan
        return None
    return None


def files_with_run_date(files, days, window=None):
    """Outputs that contain (or are named with) the date or clock time the check ran at -- probably a
    run timestamp, which two back-to-back runs can share (same minute / same day) without the
    comparison noticing."""
    rx, hits = _date_regex(days, window), []
    for rel, p in files.items():
        if rx.search(rel):
            hits.append(rel)
            continue
        text = _scan_text(rel, p)
        if text and rx.search(text):
            hits.append(rel)
    return hits


# ─── the check folders (.reproduce-check/) ─────────────────────────────────

def _write_marker(keep, status, **extra):
    try:
        data = {"tool": TOOL, "status": status, "pid": os.getpid(),
                "time": _dt.datetime.now().isoformat(timespec="seconds"), "started": _started(keep), **extra}
        text = json.dumps(data, ensure_ascii=False)
        try:
            text.encode("utf-8")
        except UnicodeEncodeError:                     # a file name that is not valid UTF-8: escape it
            text = json.dumps(data)
        (keep / MARKER).write_text(text, encoding="utf-8")
    except OSError:
        pass


def _read_marker(folder):
    try:
        data = json.loads((folder / MARKER).read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def _made_here(marker):
    """True if a status file was written by this tool (earlier versions wrote no "tool" key)."""
    if marker.get("tool") == TOOL:
        return True
    return marker.get("status") in ("running", "finished", "failed", "interrupted") and \
        isinstance(marker.get("pid"), int) and isinstance(marker.get("started"), (int, float))


def _started(folder):
    """When a check started: from its status file, else from the folder name."""
    t = _read_marker(folder).get("started")
    if isinstance(t, (int, float)):
        return t
    try:
        return _dt.datetime.strptime(folder.name[:15], "%Y%m%d-%H%M%S").timestamp()
    except ValueError:
        return 0.0


def _named_like_checks(keep_root):
    try:
        return [p for p in keep_root.iterdir() if _is_real_dir(p) and CHECK_DIR_RE.match(p.name)]
    except OSError:
        return []


def _check_folders(keep_root):
    """[(folder, status)] of the checks under keep_root that have a status file, oldest first."""
    out = [(p, _read_marker(p).get("status") or "unknown")
           for p in _named_like_checks(keep_root) if (p / MARKER).is_file()]
    return sorted(out, key=lambda x: (_started(x[0]), x[0].name))


def _unmarked(keep_root):
    """Folders that look like checks of an older version (named like one, holding only what a check
    writes) but have no status file. They are never deleted automatically."""
    out = []
    for p in _named_like_checks(keep_root):
        if (p / MARKER).exists():
            continue
        try:
            if all(CHECK_PARTS_RE.match(c.name) for c in p.iterdir()):
                out.append(p.name)
        except OSError:
            continue
    return sorted(out)


def _prunable(folder, cwd, copy_dirs, memo):
    """-> None if --keep-last may delete this check folder, else the reason why not. Only folders
    this tool made, holding only what a check writes, and whose pre-existing/ files were all either
    written by an earlier check's run (the status file records them) or have their content somewhere
    else: in their original place (written again by the runs with the same content, or put back), or
    in copy_dirs (run folders and pre-existing/ of the checks that are kept)."""
    marker = _read_marker(folder)
    if not _made_here(marker):
        return "没有可识别的本工具状态文件，不是本工具建的文件夹"
    try:
        foreign = sorted(c.name for c in folder.iterdir() if not CHECK_PARTS_RE.match(c.name))
    except OSError as e:
        return f"无法读取（{e}）"
    if foreign:
        return f"里面有不是本工具写的文件（{'、'.join(foreign[:3])}）"
    base = Path(marker["cwd"]) if isinstance(marker.get("cwd"), str) and os.path.isdir(marker["cwd"]) else cwd
    record = marker.get(FROM_CHECKS) if isinstance(marker.get(FROM_CHECKS), dict) else None
    gone, changed = _only_copies(folder / "pre-existing", base, copy_dirs, memo, record)
    if not (gone or changed):
        return None
    where = []
    if gone:
        where.append(f"{len(gone)} 个在原处已经没有了（{_names(gone, 3)}）")
    if changed:
        where.append(f"{len(changed)} 个和原处现在的同名文件内容不同（{_names(changed, 3)}）")
    if record is None:
        origin = ["这个文件夹的状态文件没有记录其中哪些是较早的检查运行生成的（较早版本的本工具建的，或那次检查"
                  "没有走到记录这一步），只能当作可能是手工做的或改过的版本"]
    else:
        edited = [rel for rel in gone + changed if rel in record]
        others = [rel for rel in gone + changed if rel not in record]
        origin = []
        if others:
            origin.append((("它" if len(others) == 1 else "它们") if not edited else _names(others, 3) + " ") +
                          "在那次检查开始时和当时留着的各次检查的运行结果都不同（手工做的或改过的，或是在检查之外运行分析得到的）")
        if edited:
            origin.append(f"{_names(edited, 3)} 是较早的检查运行生成的，但这里的副本在那次检查之后被改过")
    return (f"pre-existing/ 里检查前就有的文件中，有 {'，'.join(where)}；{'；'.join(origin)}；"
            "在原处和留下的检查副本里都没有内容相同的一份，那里是这些版本唯一的副本")


def _folder_size(path):
    total = 0
    for root, _, files in os.walk(path):
        for f in files:
            try:
                total += os.lstat(os.path.join(root, f)).st_size
            except OSError:
                pass
    return total


def _prune(keep_root, current, keep_last, cwd, removed):
    """Delete the check folders older than the keep_last most recent ones (never an unfinished one,
    never one _prunable refuses). Deleted names are appended to removed as they go (so an
    interruption still knows them). -> [{"folder", "why"}] kept.

    A pre-existing file an earlier check's run wrote (recorded in the status file) needs no copy; any
    other only counts as copied elsewhere when that copy survives this pass: in its place, or in a run
    folder or pre-existing/ of a check that is kept. Oldest first, so the oldest folder holding a
    version (often the user's own) is the one kept; everything is decided before anything is deleted."""
    folders = _check_folders(keep_root)
    finished = [p for p, st in folders if st != "running" or p == current]
    old = [p for p in finished[:max(0, len(finished) - keep_last)] if p != current]
    stay = [p for p, _ in folders if p not in old]
    kept, doomed, memo = [], [], {}
    for p in old:
        why = _prunable(p, cwd, _run_dirs(stay, pre_existing=True), memo)
        if why:
            kept.append({"folder": p.name, "why": why})
            stay.append(p)
        else:
            doomed.append(p)
    for p in doomed:
        try:
            shutil.rmtree(p)
            removed.append(p.name)
        except OSError as e:
            kept.append({"folder": p.name, "why": f"删除失败（{e}）"})
    return kept


# ─── report ─────────────────────────────────────────────────────────────────

CRASH_HINT = {"SIGSEGV": "程序崩溃", "SIGBUS": "程序崩溃", "SIGABRT": "程序异常中止",
              "SIGKILL": "多半是内存不足被系统结束，或被手动结束"}


def _exit_text(x):
    sig, rc = x.get("signal"), x["returncode"]
    if not sig:
        return f"退出码 {rc}"
    why = f"，{CRASH_HINT[sig]}" if sig in CRASH_HINT else ""
    if rc is not None and rc < 0:
        return f"被信号 {sig} 结束（退出码 {rc}{why}）"
    return f"退出码 {rc}（可能是被信号 {sig} 结束{why}）"


def _run_text(x):
    if x.get("launch_error"):
        return f"第 {x['run']} 次 无法启动（{x['launch_error']}）"
    state = "超时" if x["timed_out"] else _exit_text(x)
    return f"第 {x['run']} 次 {state}，{x['seconds']} 秒"


STREAMS = (("stdout", "屏幕输出（stdout）", "--compare-stdout"), ("stderr", "错误输出（stderr）", "--compare-stderr"))


def _stream_tag(label):
    return f"({label.replace('（', ' ').rstrip('）')})"          # "(屏幕输出 stdout)"


def _file_set_note(files, cmp):
    """Where the report shows that these files (an earlier run wrote them, the last run did not)
    differ between runs: among the differences above, or nowhere, because --exclude left them out."""
    if not cmp:
        return ""
    excluded = set(cmp.get("excluded") or ())
    listed = {d["file"] for d in cmp.get("different") or ()}
    out = [f for f in files if f in excluded]
    note = ""
    if out:
        who = ("它" if len(files) == 1 else "它们") if len(out) == len(files) else f"其中 {_names(out, 5)} "
        note = (f"；{who}被 --exclude 排除在比较之外：有的运行生成了、最后一次没有生成，这个差异没有计入结论，"
                "上面也没有列出")
    if any(f in listed for f in files):
        note += "；" + ("其余的" if out else "") + "文件集合的差异见上"
    return note


def render(result):
    r = result
    verdict = {0: "一致", 1: "不一致", 2: "未完成（运行失败、被中断或没有可比较的输出）"}[r["exit_code"]]
    lines = ["# 重跑一致性检查", "",
             f"- 结论：**{verdict}**",
             f"- 命令：`{r['command']}`（工作目录 {r['cwd']}；每次都在新的子进程里运行，"
             "运行前清空 --outputs，工作目录里的其他文件不动）"]
    if r["runs"]:
        lines.append("- 各次运行：" + "；".join(_run_text(x) for x in r["runs"]))
    cmp = r.get("comparison")
    if cmp:
        lines.append(f"- 比较的文件 {cmp['n_files']} 个：完全相同 {len(cmp['identical'])}，"
                     f"在容差内相同 {len(cmp['within_tolerance'])}（rtol={r['rtol']:g}，atol={r['atol']:g}；"
                     f"整数须完全相同），不同 {len(cmp['different'])}")
        if cmp.get("excluded"):
            lines.append(f"- 没有比较（--exclude）：{'、'.join(cmp['excluded'][:20])}"
                         + (f" 等 {len(cmp['excluded'])} 个" if len(cmp["excluded"]) > 20 else ""))
        for key, label, flag in STREAMS:
            so = cmp.get(key)
            if not so or (key == "stderr" and so["empty"] and not so["compared"]):
                continue
            if so["empty"]:
                state = "没有打印内容"
            elif so["equal"]:
                state = "各次相同"
            else:
                state = f"第 {so['runs'][0]} 次与第 {so['runs'][1]} 次不同"
            lines.append(f"- {label}：{state}" + ("" if so["compared"] else f"（未计入结论；加 {flag} 计入）"))
    for msg in r["messages"]:
        lines.append(f"- {msg}")
    if r.get("warnings"):
        lines += ["", "## 需要留意", ""]
        lines += [f"- {w}" for w in r["warnings"]]
    diff_items = list(cmp["different"]) if cmp else []
    for key, label, _ in STREAMS:
        so = (cmp or {}).get(key)
        if so and so["compared"] and not so["equal"]:
            diff_items.append({"file": _stream_tag(label), **so})
    if diff_items:
        lines += ["", "## 不一致的文件", ""]
        for d in diff_items:
            runs = f"第 {d['runs'][0]} 次 vs 第 {d['runs'][1]} 次，" if d.get("runs") else ""
            lines.append(f"### {d['file']}（{runs}{d['method']}）")
            lines += [f"- {x}" for x in d["details"]] or ["- （无更多细节）"]
            lines.append("")
    if cmp and cmp["within_tolerance"]:
        lines += ["", "## 在容差内相同的文件（数字末位、内嵌时间或随机编号不同）", ""]
        for d in cmp["within_tolerance"]:
            lines.append(f"- {d['file']}（{d['method']}）")
            lines += [f"  - {x}" for x in d.get("details", [])[:2]]
    live = "、".join(r["outputs"])
    lines += ["", "## 文件在哪里", ""]
    if r.get("last_run_in_place"):
        lines.append(f"- 最后一次运行的结果留在原位：{live}")
    elif r.get("restored"):
        lines.append(f"- 没有完成：已把 {r['restored']} 里的输出复制回原位（{live}）")
    else:
        lines.append(f"- 没有完成：原位（{live}）没有保留输出")
    lines.append(f"- 各次运行的副本与日志：{r['keep_dir']}（{'、'.join(r['snapshots'] + ['logs/'])}）")
    if r.get("pre_existing"):
        n_from = r.get("pre_existing_from_checks") or 0
        lines.append(f"- 检查前已有的输出已移到（没有删除）：{r['keep_dir']}/pre-existing/（{r['pre_existing']} 个文件"
                     + (f"，其中 {n_from} 个是较早一次检查运行生成的结果" if n_from else "") + "）")
    if r.get("put_back"):
        lines.append(f"- 其中没有哪次运行重新生成的 {len(r['put_back'])} 个文件已复制回原处（不是这次运行的结果）："
                     + _names(r["put_back"]))
    if r.get("pre_existing_replaced"):
        rp = r["pre_existing_replaced"]
        lines.append(f"- 其中 {len(rp)} 个文件这次运行重新生成了，但检查前的版本内容不同、也不是哪次检查运行生成的"
                     f"（手工改过，或是在检查之外运行分析得到的）：{_names(rp)}；检查前的版本只在 "
                     f"{r['keep_dir']}/pre-existing/ 里，需要的话从那里取回，--keep-last 不会删除这个文件夹")
    nb = r.get("not_put_back") or []
    reasons = {"earlier_run": "较早的运行生成过、最后一次运行没有生成：放回旧版本会混在这次的结果里",
               "dir": "那里现在是这次运行生成的目录"}
    for why in ("earlier_run", "dir", "error"):
        group = [x for x in nb if x["why"] == why or (why == "error" and x["why"] not in reasons)]
        if not group:
            continue
        text = reasons.get(why) or "复制时出错：" + "；".join(sorted({x["why"] for x in group}))[:300]
        if why == "earlier_run":
            text += _file_set_note([x["file"] for x in group], cmp)
        only = [x["file"] for x in group if x.get("only_copy", True)]
        line = f"- 检查前已有的 {len(group)} 个文件没有放回原处（{text}）：{_names([x['file'] for x in group])}"
        where = f"检查前的版本只在 {r['keep_dir']}/pre-existing/ 里，--keep-last 不会删除这个文件夹"
        if len(only) == len(group):
            line += f"；{where}"
        elif only:
            line += f"；其中 {_names(only, 5)} 的{where}，其余的检查前版本和某次检查运行生成的结果相同"
        else:
            line += "；它们检查前的版本和某次检查运行生成的结果相同"
        lines.append(line)
    kr = r.get("keep_root") or {}
    if kr:
        line = f"- {kr['path']} 里现有 {kr['n_checks']} 次检查的副本，共 {kr['mb']:.1f} MB"
        if kr.get("pruned"):
            line += (f"（按 --keep-last 删除了较早的 {len(kr['pruned'])} 次：只删本工具建的检查文件夹，而且它们 "
                     "pre-existing/ 里的每个文件要么是更早一次检查运行生成的结果，要么在原处或留下的检查副本里"
                     "还有内容相同的一份）")
        line += ("；副本里可能有患者级结果：看完可以删除，或下次加 --keep-last 1 只保留最近一次"
                 + ("；该目录有自己的 .gitignore，不会被 git 提交" if kr.get("gitignored") else
                    "；请把它加进 .gitignore，不要提交"))
        lines.append(line)
        for k in kr.get("kept") or ():
            lines.append(f"- 没有按 --keep-last 删除 {kr['path']}/{k['folder']}：{k['why']}；请先确认，再手动处理")
        if kr.get("unmarked"):
            names = kr["unmarked"]
            lines.append(f"- {kr['path']} 里另有 {len(names)} 个文件夹名字像检查副本、但没有本工具的状态文件"
                         f"（可能是旧版本留下的），不计数，也不会被自动删除：{'、'.join(names[:5])}")
    if r["exit_code"] == 1:
        lines += ["", "## 常见原因（逐条排查后重跑本检查）", ""]
        lines += [f"{i}. {h}" for i, h in enumerate(HINTS, 1)]
    return "\n".join(lines) + "\n"


# ─── CLI ────────────────────────────────────────────────────────────────────

class _UsageError(Exception):
    pass


class _Parser(argparse.ArgumentParser):
    def error(self, message):              # one line on stderr and exit 2 (main), not usage + traceback
        raise _UsageError(message)


def _parser():
    p = _Parser(
        prog="reproduce_check.py",
        description="在全新子进程里把分析命令运行 N 次（每次运行前清空 --outputs；--outputs 之外的文件，如缓存，"
                    "不清空，报告会列出能认出的缓存），比较输出文件是否一致。"
                    "退出码：0 一致；1 不一致；2 运行失败、被中断、没有可比较的输出或参数有误。",
        epilog='例：python3 reproduce_check.py --cmd "python3 analysis_script.py" --outputs results/ '
               'tables/results.csv --compare-stdout；或把命令放在 -- 之后：--outputs results/ -- Rscript analysis.R')
    p.add_argument("--cmd", help="要运行的命令（整条字符串，经 shell 执行）；也可以把命令写在 -- 之后")
    p.add_argument("--outputs", nargs="+", required=True, help="要比较的输出文件或目录（相对于 --cwd），可多个")
    p.add_argument("--runs", type=int, default=2, help="运行次数（≥2，默认 2）")
    p.add_argument("--cwd", default=".", help="运行命令的工作目录（默认当前目录）")
    p.add_argument("--timeout", type=float, default=600, help="每次运行的超时秒数（默认 600）")
    p.add_argument("--rtol", type=float, default=1e-12,
                   help="非整数的相对容差（默认 1e-12；整数、带前导零的编号和运行时的时间戳一律要求完全相同）")
    p.add_argument("--atol", type=float, default=1e-12, help="浮点绝对容差（默认 1e-12，用于接近 0 的数）")
    p.add_argument("--exclude", action="append", default=[], metavar="PATTERN",
                   help="不参与比较的输出文件，可重复：shell 通配符，相对于 --cwd（* 和 ? 不跨目录，** 跨任意层目录，"
                        "如 'results/**/*.pdf'）；不含 / 时只比文件名（'*.pdf' 排除所有 PDF）；写目录则排除其中所有文件")
    p.add_argument("--keep-dir", help="存放各次运行副本的目录（相对于 --cwd；默认 <cwd>/.reproduce-check/<时间>/）")
    p.add_argument("--keep-last", type=int, metavar="N",
                   help="检查结束后只保留最近 N 次检查的副本，删除更早的（默认全部保留）；只删本工具建的文件夹，"
                        "pre-existing/ 里有文件的这份内容在原处和留下的检查副本里都没有时，那次检查不删"
                        "（手工做的或改过的版本不会被删）；较早一次检查运行生成的结果不算（检查开始时和那次检查"
                        "最后一次运行的副本相同，状态文件里有记录）")
    p.add_argument("--compare-stdout", action="store_true",
                   help="把各次运行打印到屏幕的内容（stdout）也计入结论（结果只打印在屏幕上时必须加）")
    p.add_argument("--compare-stderr", action="store_true",
                   help="把各次运行的错误输出（stderr）也计入结论（结果用 R 的 message()、Python 的 logging 输出时要加）")
    p.add_argument("--report", help="另存 Markdown 报告（相对于 --cwd），如 reproduce-check.md")
    p.add_argument("--json", dest="json_path", help="另存 JSON 结果（相对于 --cwd）")
    return p


def _fail(msg):
    print(f"[重跑检查] 错误：{msg}", file=sys.stderr)
    return 2


def _under(cwd, raw):
    p = Path(raw).expanduser()
    return (p if p.is_absolute() else cwd / p).resolve()


def _command_tokens(cmd_text, cmd_list, shell):
    """-> [(token, is_output)]: the command's words; redirect and tee targets count as outputs."""
    if not shell:
        return [(t, False) for t in cmd_list]
    try:
        lx = shlex.shlex(cmd_text, posix=os.name == "posix", punctuation_chars=True)
        lx.whitespace_split = True
        words = list(lx)
    except ValueError:                      # unbalanced quotes: let the shell report it
        return []
    out, next_out, tee = [], False, False
    for w in words:
        if os.name != "posix" and len(w) > 1 and w[0] == w[-1] and w[0] in "\"'":
            w = w[1:-1]
        if w and set(w) <= set("<>&|;()"):
            next_out = ">" in w
            tee = tee and not ({"|", "&&", "||", ";", "&"} & {w})
            continue
        if w == "tee":
            tee = True
            continue
        out.append((w, next_out or (tee and not w.startswith("-"))))
        next_out = False
        if w.startswith("-") and "=" in w:              # --out=results/table.csv
            out.append((w.split("=", 1)[1], False))
    return out


def _check_command_paths(tokens, cwd, outputs):
    """-> (error or None, [tokens inside an output path that may be inputs], [dirs to create]).

    A script inside an output path is refused (it would be moved away before run 1). An output
    directory passed as an argument, or a redirect target inside it, is fine -- also on the next
    check, when the previous run's files are there. The directory of a redirect target inside an
    output directory is created, because the shell opens that file before the script runs."""
    named, mkdirs = [], []
    for tok, is_output in tokens:
        if not tok or len(tok) > 1024 or "\n" in tok or "\0" in tok:
            continue
        try:
            p = Path(os.path.realpath(os.path.join(str(cwd), os.path.expanduser(tok))))
        except (OSError, ValueError):
            continue
        hit = next((o for o in outputs if p == o.path or o.path in p.parents), None)
        if hit is None:
            continue
        if is_output:
            if p != hit.path and not p.parent.exists():
                mkdirs.append(p.parent)
            continue
        try:
            is_dir, is_file = p.is_dir(), p.is_file()
        except OSError:
            continue
        if is_dir:
            continue
        if is_file and p.suffix.lower() in SCRIPT_EXTS:
            return (f"命令要运行的脚本 {tok} 在输出路径 {hit.raw} 里：检查开始前输出路径里的文件会被移走，"
                    "脚本就找不到了。请把脚本放到输出目录之外，或把 --outputs 写得更具体"), [], []
        named.append(tok)
    return None, named, mkdirs


def _check_report_paths(paths, cwd, outputs, tokens):
    seen = {}
    for flag, p in paths:
        if p is None:
            continue
        if p.suffix.lower() in NOT_A_REPORT_EXTS:
            return f"{flag} {p.name} 的扩展名像脚本或数据文件；请用 .md / .json，以免覆盖它"
        key = str(p).lower()                       # the same file on a case-insensitive disk
        if key in seen:
            return f"{flag} 和 {seen[key]} 是同一个文件"
        seen[key] = flag
        if any(p == o.path or o.path in p.parents for o in outputs):
            return f"{flag} {p} 在输出路径里：下次检查会把它当成输出移走；请写到输出目录之外"
        if not p.parent.is_dir():
            return f"{flag} 的目录不存在：{p.parent}"
        if p.is_dir():
            return f"{flag} {p} 是一个目录"
        for tok, _ in tokens:
            try:
                if Path(os.path.realpath(os.path.join(str(cwd), tok))) == p:
                    return f"{flag} {p.name} 是命令里用到的文件，不能被报告覆盖"
            except (OSError, ValueError):
                continue
    return None


def _put_back(outputs, keep, st, tag, result, cwd):
    """After a failed / interrupted run: move its partial outputs aside and copy the earlier ones back."""
    k = st["k"]
    with _Deferred():
        _move_out(outputs, keep / f"run{k}-{tag}", st["keep_dirs"])
        _remove_new_dirs(st["new_dirs"])
        result["snapshots"].append(f"run{k}-{tag}/")
        source = keep / "pre-existing" if st["pre"] else (st["snapshots"][-1] if st["snapshots"] else None)
        if source is not None:
            _copy_back(outputs, source, cwd)
            result["restored"] = str(source)


def _input_hint(keep, stderr):
    """Run 1 failed and files had been moved out of --outputs: say so, naming the ones the error mentions."""
    moved = list(_files_at(keep / "pre-existing"))
    if not moved:
        return None
    hit = [rel for rel in moved if rel.rsplit("/", 1)[-1] in stderr]
    if hit:
        return (f"错误输出里提到了 {'、'.join(hit[:5])}：检查开始前它们在 --outputs 里，运行前被移到了 "
                f"{keep}/pre-existing/，运行时不在原处。如果它们是输入文件（数据、配置、上一步的结果），"
                "请移到输出目录之外再查")
    return (f"检查开始前 --outputs 里原有的 {len(moved)} 个文件（如 {'、'.join(moved[:3])}）运行前被移到了 "
            "pre-existing/，运行时不在原处：如果分析要读其中的文件（输入数据、配置），就会失败——请把输入文件移到输出目录之外")


def _names(items, n=10):
    return "、".join(items[:n]) + (" 等" if len(items) > n else "")


def _outside_outputs(before, after, result):
    """Warnings about state outside --outputs that the runs share: files they created or changed
    (left for the next run), and files named like a cache that were there before the check (every
    run reads the same copy). A cache with another name that predates the check cannot be seen."""
    fb, fa = before["files"], after["files"]
    changed = None                                       # None: there were too many files to compare
    if fb is not None and fa is not None:
        changed = sorted(rel for rel, v in fa.items() if fb.get(rel, (None, None))[:2] != v[:2])
        if changed:
            result["warnings"].append(
                f"运行中在 --outputs 之外新建或改动了 {len(changed)} 个文件：{_names(changed)}。"
                "每次运行前只清空 --outputs，这些文件会留给下一次运行："
                "如果它们是缓存或中间结果（如 cache/、_targets/、*.pkl），后面的运行可能直接读了前一次的结果，"
                "“一致”不代表能从头复现——请把它们加进 --outputs（或关掉缓存）再查")
    else:
        result["warnings"].append(
            f"--cwd 里的文件太多（超过 {TREE_MAX} 个），没有逐个比较运行前后 --outputs 之外的文件，"
            "运行中新建或改动的缓存、中间结果查不出来"
            + ("，只按名字找了像缓存的文件" if before["complete"] and after["complete"] else
               f"；名字像缓存的文件也只找了前 {TREE_WALK_MAX} 个文件")
            + "：请自己确认分析没有用缓存（有的话把缓存目录加进 --outputs）")
        new = sorted(set(after["cache_like"]) - set(before["cache_like"]))
        if new:
            result["warnings"].append(
                f"运行中在 --outputs 之外新建了 {len(new)} 个名字像缓存的文件：{_names(new)}。它们会留给下一次运行，"
                "后面的运行可能直接读了前一次的结果——请把它们加进 --outputs（或关掉缓存）再查")
    result["changed_outside_outputs"] = changed
    now = set(after["cache_like"]) - set(changed or ())
    old = [rel for rel in before["cache_like"] if rel in now]
    result["cache_like_outside_outputs"] = old[:200]
    if old:
        read = {rel for rel in old if fb is not None and fa is not None and rel in fb and rel in fa
                and fb[rel][2] != fa[rel][2]}
        shown = [rel + ("（运行中读过）" if rel in read else "") for rel in old]
        result["warnings"].append(
            f"--outputs 之外有 {len(old)} 个检查前就有、名字像缓存的文件：{_names(shown)}。每次运行前只清空 "
            "--outputs，这些文件不动，各次运行读到的都是同一份：如果分析从缓存里读结果（上一次的拟合、中间数据），"
            "各次结果自然相同，“一致”不代表能从头复现——请把缓存目录加进 --outputs（检查前会被移走，每次运行前清空）"
            "或关掉缓存后重查；确认不是缓存（如输入数据）就不用管")


def _check(args, cmd, shell, cwd, outputs, keep, result, st, named):
    for d in st["mkdirs"]:                    # "> results/stdout.txt" needs results/ before the shell starts
        try:
            d.mkdir(parents=True, exist_ok=True)
        except OSError as e:
            result["messages"].append(f"无法创建命令里重定向输出要用的目录 {d}：{e}")
            result["exit_code"] = 2
            return
    st["phase"] = "pre"
    with _Deferred():
        st["pre"] = result["pre_existing"] = _move_out(outputs, keep / "pre-existing")
        st["keep_dirs"] = _dirs_under(outputs)
    try:
        st["from_checks"] = _from_earlier_checks(keep / "pre-existing", keep, st["keep_root"], st["memo"])
        result["pre_existing_from_checks"] = len(st["from_checks"])
    except OSError:                           # not recorded: --keep-last then protects every pre-existing file
        st["from_checks"] = None
    st["new_dirs"] = _missing_parents(outputs, cwd)
    skip = {o.path for o in outputs} | {st["keep_root"]}
    before = _tree_state(cwd, skip)
    st["phase"] = "run"
    days = {_dt.date.today()}
    t_start = time.time()
    background = []
    for k in range(1, args.runs + 1):
        st["k"] = k
        res = run_once(cmd, shell, cwd, args.timeout, keep / "logs", k)
        result["runs"].append(dict({x: res[x] for x in ("run", "returncode", "timed_out", "seconds", "launch_error")},
                                   signal=_signal_of(res["returncode"], shell), background_killed=res["background"]))
        if res["background"]:
            background.append(str(k))
        if not res["ok"]:
            _put_back(outputs, keep, st, "failed", result, cwd)
            if res["launch_error"]:
                why = f"命令无法启动（{res['launch_error']}）"
            elif res["timed_out"]:
                why = f"超时（>{args.timeout:g} 秒），命令及其子进程已被结束"
            else:
                why = _exit_text(result["runs"][-1])
            tail = res["stderr"].strip().splitlines()[-8:] if not res["launch_error"] else []
            result["messages"].append(f"第 {k} 次运行失败：{why}" + ("；错误输出最后几行见下" if tail else ""))
            result["messages"] += [f"`{ln}`" for ln in tail]
            if named:
                result["messages"].append(
                    f"命令里写到的 {'、'.join(named)} 在 --outputs 路径里：检查前这些路径里原有的文件被移走、"
                    "每次运行前都会清空。如果其中有输入文件（数据、配置），请移到输出目录之外")
            elif k == 1 and st["pre"] and not res["launch_error"]:
                hint = _input_hint(keep, res["stderr"])
                if hint:
                    result["messages"].append(hint)
            result["exit_code"] = 2
            return
        snap = keep / f"run{k}"
        with _Deferred():
            if k < args.runs:
                _move_out(outputs, snap, st["keep_dirs"])
                _remove_new_dirs(st["new_dirs"])
            else:
                _copy_out(outputs, snap)
        st["snapshots"].append(snap)
        result["snapshots"].append(f"run{k}/")

    t_end = time.time()
    after = _tree_state(cwd, skip)
    st["phase"] = "restore"
    with _Deferred():
        result["put_back"], result["not_put_back"] = _restore_missing(keep / "pre-existing", cwd, st["snapshots"])
    st["phase"] = "compare"
    # files whose pre-check version is now only in pre-existing/: gone (not put back) or replaced by
    # other content -- and not something a check's run wrote (edited by hand, or made outside a check)
    run_dirs = _run_dirs(p for p, _ in _check_folders(st["keep_root"]))
    gone, result["pre_existing_replaced"] = _only_copies(keep / "pre-existing", cwd, run_dirs, st["memo"],
                                                         st["from_checks"])
    for x in result["not_put_back"]:
        x["only_copy"] = x["file"] in gone
    result["last_run_in_place"] = True
    days.add(_dt.date.today())
    _WINDOW.update(start=t_start, end=t_end)
    cmp = compare_runs(st["snapshots"], outputs, args.rtol, args.atol, args.exclude)
    streams = {}
    for key, label, flag in STREAMS:
        logs = [keep / "logs" / f"run{k}.{key}.txt" for k in range(1, args.runs + 1)]
        compared = args.compare_stdout if key == "stdout" else args.compare_stderr
        cmp[key] = streams[key] = compare_stream(logs, args.rtol, args.atol, compared)
        so = streams[key]
        if not so["compared"] and not so["equal"]:
            shown = "；".join(f"`{ln[:120]}`" for ln in so["details"][:4])
            where = ("（如 `[SAP 4.1] …` 行），请加 --compare-stdout 重跑" if key == "stdout" else
                     "（R 的 message()、Python 的 logging），请加 --compare-stderr 重跑")
            result["warnings"].append(f"{label}各次运行不同，但没有计入结论（没加 {flag}）：如果有结果只写在这里"
                                      f"{where}。不同的行：{shown}")
    result["comparison"] = cmp
    if background:
        result["warnings"].append(
            f"第 {'、'.join(background)} 次运行结束时，它启动的后台进程还在运行，已被结束（否则会写进下一次运行的输出）："
            "分析脚本应等子进程结束再退出")
    _outside_outputs(before, after, result)
    for pat in cmp["unmatched_exclude"]:
        result["warnings"].append(f"--exclude {pat!r} 没有匹配到任何输出文件（* 不跨目录；跨目录用 **/，如 'results/**/*.pdf'）")
    live = {k: v for k, v in _collect(outputs, base=st["snapshots"][-1]).items()
            if not _excluded(k, args.exclude)}
    window = (t_start, t_end)
    dated = files_with_run_date(live, days, window)
    rx = _date_regex(days, window)
    for key, label, _ in STREAMS:
        if streams[key]["compared"] and rx.search(_read_log(keep / "logs" / f"run{args.runs}.{key}.txt")):
            dated.append(_stream_tag(label))
    if dated:
        result["warnings"].append(
            f"输出里出现了运行当天的日期或运行时的钟点：{'、'.join(dated[:10])}。如果那是运行时间，连续两次运行可能恰好相同"
            "（同一分钟、同一天）而查不出来——请确认；结果文件里不要写运行时间")
    if cmp["never_produced"]:
        result["messages"].append("这些输出路径在各次运行中都没有文件，无法比较（检查 --outputs 是否写对、"
                                  "路径是否相对于 --cwd）：" + "、".join(cmp["never_produced"]))
        missing = [o for o in outputs if o.raw in cmp["never_produced"]]
        if any(r == o.rel or r.startswith(o.rel + "/") for r in result["put_back"] for o in missing):
            result["restored"] = str(keep / "pre-existing")
            result["messages"].append("这些路径原有的文件已放回原处（运行前被移到了 pre-existing/）")
        result["exit_code"] = 2
    elif cmp["n_files"] == 0:
        result["messages"].append("没有找到任何可比较的输出文件" + ("（其余都被 --exclude 排除了）" if cmp["excluded"] else ""))
        result["exit_code"] = 2
    else:
        stream_differs = any(s["compared"] and not s["equal"] for s in streams.values())
        result["exit_code"] = 1 if cmp["different"] or stream_differs else 0


def main(argv=None):
    try:
        return _main(argv)
    except (KeyboardInterrupt, Interrupted):
        if _SIG.get("keep") == "done":        # only possible in a split second around the report
            return _fail("已中断（报告已经写完）")
        if _SIG.get("keep"):
            return _fail(f"已中断，没有写完报告；检查前移走的文件和各次运行的副本在 {_SIG['keep']}/ 里")
        return _fail("已中断（检查还没有开始移动任何文件）")
    except Exception as e:                    # never a traceback, never exit 1 (that means "outputs differ")
        return _fail(f"意外错误（{type(e).__name__}: {e}）")
    finally:
        _WINDOW.update(start=None, end=None)
        _SIG["keep"] = None


def _main(argv):
    argv = list(sys.argv[1:] if argv is None else argv)
    cmd_list = None
    if "--" in argv:
        i = argv.index("--")
        argv, cmd_list = argv[:i], argv[i + 1:]
    try:
        args = _parser().parse_args(argv)
    except _UsageError as e:
        return _fail(f"{e}（python3 reproduce_check.py --help 查看用法）")
    if bool(args.cmd) == bool(cmd_list):
        return _fail("请用 --cmd \"<命令>\" 或在 -- 之后写命令（二选一）")
    if args.runs < 2:
        return _fail("--runs 至少为 2")
    if not args.timeout > 0:
        return _fail("--timeout 必须大于 0")
    if not (args.rtol >= 0 and args.atol >= 0):
        return _fail("--rtol / --atol 不能是负数")
    if args.keep_last is not None and args.keep_last < 1:
        return _fail("--keep-last 至少为 1")
    cwd = Path(args.cwd).expanduser().resolve()
    if not cwd.is_dir():
        return _fail(f"工作目录不存在：{cwd}")
    shell = bool(args.cmd)
    cmd = args.cmd if shell else cmd_list
    cmd_text = args.cmd if shell else " ".join(shlex.quote(c) for c in cmd_list)
    args.exclude = [_norm_pattern(p, cwd) for p in args.exclude]

    outputs = []
    for raw in args.outputs:
        try:
            outputs.append(Output(raw, cwd))
        except (OSError, RuntimeError, ValueError) as e:     # a symlink loop, an unreadable path
            return _fail(f"无法解析 --outputs {raw}：{e}")
    for o in outputs:
        if o.path == cwd or o.path in cwd.parents:
            return _fail(f"--outputs {o.raw} 是工作目录本身或它的上级目录；只能列出分析生成的文件或目录")
        if cwd not in o.path.parents:
            # outside the project the checker would move files it has no business touching
            return _fail(f"--outputs {o.raw} 不在工作目录 {cwd} 里；只能列出这次分析在项目里生成的文件或目录")
    for a in outputs:
        for b in outputs:
            if a is not b and (a.path == b.path or b.path in a.path.parents):
                return _fail(f"--outputs 有重叠：{a.raw} 在 {b.raw} 里面")
    tokens = _command_tokens(cmd_text, cmd_list, shell)
    err, named, mkdirs = _check_command_paths(tokens, cwd, outputs)
    if err:
        return _fail(err)
    report = _under(cwd, args.report) if args.report else None
    json_path = _under(cwd, args.json_path) if args.json_path else None
    err = _check_report_paths((("--report", report), ("--json", json_path)), cwd, outputs, tokens)
    if err:
        return _fail(err)

    keep_root = _under(cwd, args.keep_dir) if args.keep_dir else cwd / ".reproduce-check"
    for o in outputs:
        if keep_root == o.path or o.path in keep_root.parents or keep_root in o.path.parents:
            return _fail("--keep-dir 不能和输出路径重叠")
    new_root = not keep_root.exists()
    try:
        keep_root.mkdir(parents=True, exist_ok=True)
        gitignored = (keep_root / ".gitignore").is_file()
        if not gitignored and (new_root or not args.keep_dir):
            (keep_root / ".gitignore").write_text(GITIGNORE, encoding="utf-8")
            gitignored = True
        keep = Path(tempfile.mkdtemp(prefix=_dt.datetime.now().strftime("%Y%m%d-%H%M%S-"), dir=str(keep_root)))
        (keep / MARKER).write_text(json.dumps({"tool": TOOL, "started": time.time()}), encoding="utf-8")
    except OSError as e:
        return _fail(f"无法创建副本目录 {keep_root}：{e}")
    stale = [p for p, status in _check_folders(keep_root) if status == "running" and p != keep]
    _write_marker(keep, "running", command=cmd_text, cwd=str(cwd))

    result = {"command": cmd_text, "cwd": str(cwd), "outputs": [o.raw for o in outputs], "runs": [],
              "n_runs": args.runs, "rtol": args.rtol, "atol": args.atol, "keep_dir": str(keep),
              "snapshots": [], "messages": [], "warnings": [], "pre_existing": 0, "pre_existing_from_checks": 0,
              "restored": None,
              "put_back": [], "not_put_back": [], "interrupted": False, "last_run_in_place": False,
              "comparison": None, "keep_root": None, "exit_code": 2}
    for p in stale:
        if _files_at(p / "pre-existing"):
            result["warnings"].append(f"较早的一次检查（{p}）没有正常结束（可能被强行终止）：它运行前移走的文件"
                                      f"还在 {p}/pre-existing/，请确认是否需要放回原处")
    st = {"phase": "start", "k": 0, "pre": 0, "snapshots": [], "keep_dirs": None, "new_dirs": [],
          "mkdirs": mkdirs, "keep_root": keep_root, "from_checks": None, "memo": {}}
    old = _install_handlers()                 # kept until the report is written
    _SIG["keep"] = str(keep)
    try:
        try:
            _check(args, cmd, shell, cwd, outputs, keep, result, st, named)
        except (Interrupted, KeyboardInterrupt) as e:
            name = _signal_name(e.signum) if isinstance(e, Interrupted) else "SIGINT"
            result["interrupted"] = True
            _abort(result, st, outputs, keep, cwd, f"检查被中断（{name}）：分析命令已被结束")
        except Exception as e:                # an unexpected error must not leave the outputs moved away
            _abort(result, st, outputs, keep, cwd, f"检查出错（{type(e).__name__}: {e}）")
        _housekeeping(args, result, keep, keep_root, cwd, cmd_text, gitignored, st["from_checks"])
        _SIG["defer"] += 1                    # from here on a signal only waits: the report gets written
        code = _finish(result, report, json_path)
        _SIG["keep"] = "done"
        return code
    finally:
        _restore_handlers(old)
        _SIG["defer"], _SIG["pending"] = 0, None


def _housekeeping(args, result, keep, keep_root, cwd, cmd_text, gitignored, from_checks=None):
    """Status file (with from_checks: which pre-existing/ files an earlier check's run wrote; None when
    the check stopped before that was known), --keep-last and the size of the copies. The check is
    over: neither an error nor Ctrl-C here may cost the user the report."""
    removed, pruning = [], False
    try:
        status = "interrupted" if result["interrupted"] else ("finished" if result["exit_code"] in (0, 1) else "failed")
        extra = {FROM_CHECKS: from_checks} if from_checks is not None else {}
        _write_marker(keep, status, command=cmd_text, cwd=str(cwd), exit_code=result["exit_code"], **extra)
        pruning = bool(args.keep_last)
        kept = _prune(keep_root, keep, args.keep_last, cwd, removed) if args.keep_last else []
        pruning = False
        result["keep_root"] = {"path": str(keep_root), "n_checks": len(_check_folders(keep_root)),
                               "mb": _folder_size(keep_root) / 1e6, "pruned": removed, "kept": kept,
                               "unmarked": _unmarked(keep_root), "gitignored": gitignored}
        return
    except (Interrupted, KeyboardInterrupt) as e:
        name = _signal_name(e.signum) if isinstance(e, Interrupted) else "SIGINT"
        why = f"被中断（{name}）：检查结果不受影响"
        partly = "；可能有一个较早的检查副本只删了一部分" if pruning else ""
    except Exception as e:
        why, partly = f"出错（{type(e).__name__}: {e}）", ""
    done = (f"按 --keep-last 已删除较早的检查副本 {_names(removed, 5)}（都是判定可以删的）" if removed
            else "没有删除任何较早的检查副本")
    result["messages"].append(f"整理副本目录 {keep_root} 时{why}；{done}{partly}")


def _abort(result, st, outputs, keep, cwd, message):
    _SIG["defer"] += 1                        # a second Ctrl-C must not stop the files going back
    try:
        if st["phase"] == "pre":
            if (keep / "pre-existing").exists():
                _copy_back(outputs, keep / "pre-existing", cwd)
                result["restored"] = str(keep / "pre-existing")
        elif st["phase"] == "run":
            _put_back(outputs, keep, st, "interrupted" if result["interrupted"] else "failed", result, cwd)
        elif st["phase"] == "restore":
            result["put_back"], result["not_put_back"] = _restore_missing(keep / "pre-existing", cwd, st["snapshots"])
        result["messages"].append(message)
    except Exception as e:
        result["messages"].append(message)
        result["messages"].append(f"放回原有输出时出错（{type(e).__name__}: {e}）：原有输出在 {keep}/pre-existing/，"
                                  "请手动放回")
    finally:
        _SIG["defer"] -= 1
        _SIG["pending"] = None
    result["exit_code"] = 2


def _finish(result, report, json_path):
    text = render(result)
    for path, content in ((report, text), (json_path, json.dumps(result, ensure_ascii=False, indent=2))):
        if path is None:
            continue
        try:
            path.write_text(content, encoding="utf-8")
        except OSError as e:
            print(f"[重跑检查] 错误：无法写入 {path}：{e}", file=sys.stderr)
            result["exit_code"] = 2
    print(text.rstrip())
    return result["exit_code"]


if __name__ == "__main__":
    sys.exit(main())
